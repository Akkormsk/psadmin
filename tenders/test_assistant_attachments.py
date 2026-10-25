import json
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from .assistant_agent import _validate_action_text_sources
from .assistant_conversations import (
    MAX_ATTACHMENT_TEXT_CHARS, active_attachments, create_conversation,
    message_attachments, respond, serialize_attachment, upload_attachment,
)
from .assistant_tools import AssistantToolError, available_tools
from .models import ConversationAttachment
from .test_assistant_conversations import function_call, planned_call


def document_bytes(kind, text):
    output = BytesIO()
    if kind == 'docx':
        document = Document()
        document.add_paragraph(text)
        document.save(output)
    elif kind == 'xlsx':
        book = Workbook()
        book.active.append([text, 7, 12.5])
        book.save(output)
    elif kind == 'pdf':
        writer = PdfWriter()
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 40 720 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.write(output)
    elif kind in ('png', 'scan'):
        Image.new('RGB', (8, 8), 'white').save(output, format='PDF' if kind == 'scan' else 'PNG')
    else:
        return text.encode('utf-8-sig')
    return output.getvalue()


class AssistantAttachmentTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('files', 'files@example.test', 'password')
        self.conversation = create_conversation(self.user)
        self.client.force_login(self.user)
        for target in ('tenders.assistant_agent.preflight', 'tenders.services._ai_gateway_json'):
            mocked = patch(target, side_effect=AssertionError('Paid calls forbidden') if target.endswith('_ai_gateway_json') else None)
            mocked.start()
            self.addCleanup(mocked.stop)

    def upload(self, kind, text='FILE_CONTENT'):
        return upload_attachment(self.conversation, self.user, SimpleUploadedFile('sample.' + ('pdf' if kind == 'scan' else kind), document_bytes(kind, text)))

    def test_supported_document_content_reaches_gateway(self):
        for kind in ('txt', 'csv', 'docx', 'xlsx', 'pdf'):
            with self.subTest(kind=kind):
                attachment = self.upload(kind)
                with patch('tenders.assistant_agent._gateway', return_value=function_call('assistant_chat', {'answer': 'Ответ'})) as gateway:
                    respond(self.conversation, self.user, 'Прочитай файл', attachments=[attachment])
                messages = gateway.call_args.args[0]
                self.assertTrue('FILE_CONTENT' in json.dumps(messages))
                self.assertIn(attachment.name, json.dumps(messages))

    def test_image_reaches_gateway_as_vision_content(self):
        attachment = self.upload('png')
        with patch('tenders.assistant_agent._gateway', return_value=function_call('assistant_chat', {'answer': 'Ответ'})) as gateway:
            respond(self.conversation, self.user, 'Что на картинке?', attachments=[attachment])
        parts = next(m['content'] for m in gateway.call_args.args[0] if m['role'] == 'user')
        self.assertIsInstance(parts, list)
        self.assertTrue(any(p.get('image_url', {}).get('url', '').startswith('data:image/png;base64,') for p in parts))

    def test_image_is_retained_without_copying_base64_on_clarification(self):
        attachment = self.upload('png')
        with patch('tenders.assistant_agent._gateway', side_effect=[
            planned_call('cash.transaction_create'), function_call('assistant_ask', {'question': 'Дата?'}),
            function_call('assistant_ask', {'question': 'Счёт?'}),
        ]):
            respond(self.conversation, self.user, 'Посмотри чек', attachments=[attachment])
            respond(self.conversation, self.user, 'Сегодня')
        messages = self.conversation.context['agent_run']['messages']
        self.assertEqual(json.dumps(messages).count('data:image/png;base64,'), 1)
        self.assertTrue(any(m == {'role': 'user', 'content': 'Сегодня'} for m in messages))

    def test_replacement_and_followup_keep_current_file_and_original_evidence(self):
        first, second = self.upload('txt', 'FIRST_CONTENT'), self.upload('txt', 'SECOND_CONTENT')
        replies = [planned_call('read:cash.transactions'), function_call('assistant_ask', {'question': 'Период?'}), function_call('assistant_ask', {'question': 'Счёт?'}), function_call('assistant_ask', {'question': 'Уточни?'})]
        with patch('tenders.assistant_agent._gateway', side_effect=replies):
            respond(self.conversation, self.user, 'Проверь', attachments=[first])
            respond(self.conversation, self.user, 'Исправленный файл', attachments=[second])
            respond(self.conversation, self.user, 'Продолжай')
        self.conversation.refresh_from_db()
        state = self.conversation.context['agent_run']
        self.assertEqual(state['attachments'][0]['id'], second.pk)
        self.assertEqual([a.pk for a in active_attachments(self.conversation, self.user)], [second.pk])
        self.assertIn('FIRST_CONTENT', json.dumps(state['messages']))
        self.assertEqual(json.dumps(state['messages']).count('SECOND_CONTENT'), 1)

    def test_file_text_is_grounded_but_invented_action_text_is_rejected(self):
        attachment = self.upload('txt', 'Покупка бумаги')
        with patch('tenders.assistant_agent._gateway', side_effect=[planned_call('cash.transaction_create'), function_call('assistant_ask', {'question': 'Дата?'})]):
            respond(self.conversation, self.user, 'Создай расход по файлу', attachments=[attachment])
        tool = next(t for t in available_tools(self.user) if t['id'] == 'cash.transaction_create')
        state = self.conversation.context['agent_run']
        _validate_action_text_sources(tool, {'reason': 'Покупка бумаги'}, state)
        with self.assertRaises(AssistantToolError):
            _validate_action_text_sources(tool, {'reason': 'Покупка автомобиля'}, state)

    def test_file_data_prepares_confirmation_without_executing_action(self):
        from cash.models import CashTransaction
        attachment = self.upload('txt', '2026-10-25; наличные; расход; 125; Покупка бумаги')
        arguments = {'operation_date': '2026-10-25', 'account': 'cash', 'direction': 'expense', 'amount': 125, 'reason': 'Покупка бумаги'}
        with patch('tenders.assistant_agent._gateway', side_effect=[
            planned_call('cash.transaction_create'), function_call('cash_transaction_create', arguments),
        ]):
            card = respond(self.conversation, self.user, 'Подготовь расход по файлу', attachments=[attachment])[-1]
        self.assertEqual(card['kind'], 'confirmation_required')
        self.assertEqual(self.conversation.context['agent_run']['waiting_confirmation']['arguments'], arguments)
        self.assertFalse(CashTransaction.objects.exists())

    def test_corrupt_and_unsupported_uploads_return_json_without_persisting(self):
        url = reverse('assistant_conversation_attachment', args=[self.conversation.pk])
        for name in ('bad.exe', 'bad.docx', 'bad.xlsx', 'bad.pdf'):
            with self.subTest(name=name):
                response = self.client.post(url, {'file': SimpleUploadedFile(name, b'broken')})
                self.assertEqual(response.status_code, 400)
                self.assertTrue(response.json()['detail'])
        self.assertFalse(ConversationAttachment.objects.exists())

    def test_empty_extraction_and_limit_are_visible_in_metadata_and_model_input(self):
        for attachment in (self.upload('scan'), self.upload('txt', 'A' * (MAX_ATTACHMENT_TEXT_CHARS + 1))):
            data = serialize_attachment(attachment, include_text=True)
            self.assertTrue(data['warning'])
            with patch('tenders.assistant_agent._gateway', return_value=function_call('assistant_chat', {'answer': 'Ответ'})) as gateway:
                messages = respond(self.conversation, self.user, 'Прочитай', attachments=[attachment])
            self.assertIn(data['warning'], json.dumps(gateway.call_args.args[0], ensure_ascii=False))
            self.assertEqual(messages[-2]['data']['attachments'][0]['warning'], data['warning'])

    def test_owner_conversation_limits_and_anonymous_access_are_preserved(self):
        attachment = self.upload('txt')
        other = get_user_model().objects.create_user('other-files')
        for conversation, user in ((self.conversation, other), (create_conversation(self.user), self.user)):
            with self.assertRaises(ValueError):
                message_attachments(conversation, user, [attachment.pk])
        with self.assertRaises(ValueError):
            message_attachments(self.conversation, self.user, [attachment.pk] * 6)
        self.client.logout()
        self.assertEqual(self.client.post(reverse('assistant_conversation_attachment', args=[self.conversation.pk])).status_code, 302)
