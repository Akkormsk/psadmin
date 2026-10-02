from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from tenders.models import ProcessDefinition, Proposal
from tenders.proposals import apply_proposal

COLUMNS = (
    "name", "supplies_input", "performs_production", "terminal_mode",
    "scope_tags", "when_to_use", "when_not_to_use", "parameters",
)
TERMINAL_MAP = {"всегда": "always", "никогда": "never", "иногда": "sometimes"}


def _bool(value):
    return str(value or "").strip().lower() == "да"


def _split(value, sep):
    return [part.strip() for part in str(value or "").split(sep) if part.strip()]


def _row_to_payload(row):
    name, supplies, performs, terminal, scope, when_to, when_not, required = row[:8]
    name = str(name or "").strip()
    if not name:
        return None
    return {
        "name": name,
        "supplies_input": _bool(supplies),
        "performs_production": _bool(performs),
        "terminal_mode": TERMINAL_MAP.get(str(terminal or "").strip().lower(), "sometimes"),
        "scope_tags": _split(scope, ","),
        "when_to_use": str(when_to or "").strip(),
        "when_not_to_use": str(when_not or "").strip(),
        "parameters": {"required": _split(required, ";"), "optional": []},
    }


class Command(BaseCommand):
    """Массовый импорт этапов «Базы производства» из Excel — тот же путь,
    что и ручное создание одного этапа: Proposal(type=create_stage),
    сразу applied (импорт таблицы администратором — это и есть явное
    подтверждение, см. tenders/proposals.py). Повторный импорт того же
    файла обновляет существующие по имени, не плодит дубли (см.
    _apply_create_stage: get_or_create по name)."""

    help = "Импорт этапов из Excel: Название | Даёт заготовку/материал | Выполняет производство | Завершает маршрут | Что производим | Когда использовать | Когда не использовать | Что обязательно знать из ТЗ"

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path")
        parser.add_argument("--sheet", default=None)
        parser.add_argument("--replace", action="store_true", help="Удалить текущие этапы и связи перед импортом")

    def handle(self, *args, **options):
        try:
            workbook = load_workbook(options["xlsx_path"], data_only=True)
        except FileNotFoundError:
            raise CommandError(f"Файл не найден: {options['xlsx_path']}")
        sheet = workbook[options["sheet"]] if options["sheet"] else workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))[1:]  # первая строка — заголовки

        user = get_user_model().objects.filter(is_superuser=True).first()
        if user is None:
            raise CommandError("Не найден ни один суперпользователь — не на кого оформить импорт.")

        created, updated, skipped = 0, 0, 0
        with transaction.atomic():
            if options["replace"]:
                deleted, _ = ProcessDefinition.objects.all().delete()
                self.stdout.write(f"Справочник очищен: удалено объектов вместе со связями — {deleted}.")
            for row in rows:
                payload = _row_to_payload(row)
                if payload is None:
                    skipped += 1
                    continue
                existed = ProcessDefinition.objects.filter(name=payload["name"]).exists()
                proposal = Proposal.objects.create(
                    type=Proposal.TYPE_CREATE_STAGE, payload=payload,
                    summary=f"Импорт из таблицы: «{payload['name']}»",
                    source_text="tenders_stage_catalog_import.xlsx", created_by=user,
                )
                apply_proposal(proposal, user)
                if existed:
                    updated += 1
                else:
                    created += 1
                self.stdout.write(f"{'обновлён' if existed else 'создан'}: {payload['name']}")

        self.stdout.write(self.style.SUCCESS(f"Готово: создано {created}, обновлено {updated}, пропущено {skipped}."))
