"""Предпросмотр документов извещения: качаем файл с zakupki.gov.ru у себя на сервере
и вытаскиваем текст/таблицы. Требует, чтобы сервер доставал zakupki.gov.ru (прод, РФ).
Разбор .docx/.xlsx учитывает объединённые ячейки (colspan/rowspan).
"""
from __future__ import annotations

import io
import re
import zipfile
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from django.utils.html import escape

ALLOWED_PREFIX = "https://zakupki.gov.ru/"
MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 300
MAX_PDF_PAGES = 40
MAX_PARAGRAPHS = 1200


class DocumentError(RuntimeError):
    pass


def fetch_document(url: str, *, timeout: int = 40) -> bytes:
    if not url.startswith(ALLOWED_PREFIX):
        raise DocumentError("Ссылка не с портала ЕИС.")
    request = Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
        "Accept": "*/*",
    })
    try:
        with urlopen(request, timeout=timeout) as response:
            data = response.read(MAX_BYTES + 1)
    except HTTPError as exc:
        raise DocumentError(f"Портал ЕИС ответил HTTP {exc.code}.") from exc
    except (URLError, TimeoutError, ConnectionError, OSError) as exc:
        raise DocumentError("Не удалось скачать файл с портала ЕИС (таймаут или недоступен). Откройте оригинал по ссылке.") from exc
    if len(data) > MAX_BYTES:
        raise DocumentError("Файл слишком большой для предпросмотра.")
    return data


def _ext(filename: str) -> str:
    return (re.sub(r"\.zip$", "", (filename or "").lower())).rsplit(".", 1)[-1]


def _docx_vmerge(tc) -> str | None:
    """None / 'restart' / 'continue' — вертикальное объединение ячейки таблицы."""
    from docx.oxml.ns import qn

    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return None
    vm = tc_pr.find(qn("w:vMerge"))
    if vm is None:
        return None
    return vm.get(qn("w:val")) or "continue"


def _docx_gridspan(tc) -> int:
    """Сколько столбцов сетки занимает ячейка (w:gridSpan), 1 если не задано."""
    from docx.oxml.ns import qn

    tc_pr = tc.find(qn("w:tcPr"))
    if tc_pr is None:
        return 1
    gs = tc_pr.find(qn("w:gridSpan"))
    if gs is None:
        return 1
    try:
        return int(gs.get(qn("w:val")) or 1)
    except (TypeError, ValueError):
        return 1


def _docx_cell_text(tc) -> str:
    from docx.oxml.ns import qn

    paras = []
    for p in tc.findall(qn("w:p")):
        text = "".join(node.text or "" for node in p.iter(qn("w:t"))).strip()
        if text:
            paras.append(escape(text))
    return "<br>".join(paras)


def _docx_table_html(table) -> str:
    """Учитывает объединение ячеек (colspan через w:gridSpan, rowspan через w:vMerge).

    Важно: работаем с СЫРЫМ XML (w:tr/w:tc) каждой строки, а не через table.rows[i].cells —
    python-docx для вертикально объединённых ячеек «схлопывает» row.cells так, что ячейка
    строки-продолжения возвращается как ТОТ ЖЕ объект, что и ячейка-«шапка» (с vMerge=restart),
    а не её собственный tc с vMerge=continue. Из-за этого проверка на continue никогда не
    срабатывала и объединённые по вертикали ячейки (частые в спецификациях/сметах ЕИС)
    дублировались на каждой строке вместо rowspan — таблица «съезжала».
    """
    from docx.oxml.ns import qn

    trs = table._tbl.findall(qn("w:tr"))
    rows_tcs = [tr.findall(qn("w:tc")) for tr in trs]
    n_rows = len(rows_tcs)

    # Столбец начала каждой ячейки в её строке (с учётом gridSpan предыдущих ячеек той же строки).
    col_of = []
    for tcs in rows_tcs:
        cols, c = [], 0
        for tc in tcs:
            cols.append(c)
            c += _docx_gridspan(tc)
        col_of.append(cols)

    def tc_at(r2: int, col_idx: int):
        for tc, c0 in zip(rows_tcs[r2], col_of[r2]):
            if c0 <= col_idx < c0 + _docx_gridspan(tc):
                return tc
        return None

    rows_html = []
    for r, tcs in enumerate(rows_tcs):
        cells_html = []
        for tc, col_idx in zip(tcs, col_of[r]):
            if _docx_vmerge(tc) == "continue":
                continue
            colspan = _docx_gridspan(tc)
            rowspan = 1
            if _docx_vmerge(tc) == "restart":
                for r2 in range(r + 1, n_rows):
                    below = tc_at(r2, col_idx)
                    if below is not None and _docx_vmerge(below) == "continue":
                        rowspan += 1
                    else:
                        break
            attrs = (f' colspan="{colspan}"' if colspan > 1 else "") + (f' rowspan="{rowspan}"' if rowspan > 1 else "")
            cells_html.append(f"<td{attrs}>{_docx_cell_text(tc)}</td>")
        if cells_html:
            rows_html.append("<tr>" + "".join(cells_html) + "</tr>")
    return f'<table class="ts-doc-table">{"".join(rows_html)}</table>' if rows_html else ""


def _docx_html(data: bytes) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table

    doc = Document(io.BytesIO(data))
    out, count = [], 0
    for child in doc.element.body.iterchildren():
        if count > MAX_PARAGRAPHS:
            out.append("<p>…</p>")
            break
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            if text:
                out.append(f"<p>{escape(text)}</p>")
                count += 1
        elif child.tag == qn("w:tbl"):
            html = _docx_table_html(Table(child, doc))
            if html:
                out.append(html)
                count += 1
    return "\n".join(out) or "<p class='ts-sub'>Документ без текста.</p>"


def _xlsx_html(data: bytes) -> str:
    from openpyxl import load_workbook

    # read_only не отдаёт merged_cells вообще — без него объединённые ячейки (частые
    # в сметах/спецификациях ЕИС) превращаются в «съехавшую» таблицу. Файлы уже
    # ограничены по размеру (MAX_BYTES), поэтому обычная загрузка не проблема.
    wb = load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for sheet in wb.worksheets:
        out.append(f"<h4>{escape(sheet.title)}</h4>")
        span_at = {}  # (row, col) -> (rowspan, colspan) для левой верхней ячейки объединения
        skip = set()  # (row, col) остальных ячеек объединения — не выводим повторно
        for mc in sheet.merged_cells.ranges:
            span_at[(mc.min_row, mc.min_col)] = (mc.max_row - mc.min_row + 1, mc.max_col - mc.min_col + 1)
            for rr in range(mc.min_row, mc.max_row + 1):
                for cc in range(mc.min_col, mc.max_col + 1):
                    if (rr, cc) != (mc.min_row, mc.min_col):
                        skip.add((rr, cc))

        rows, row_count = [], 0
        for row in sheet.iter_rows():
            if row_count >= MAX_ROWS:
                rows.append("<tr><td>… ещё строки</td></tr>")
                break
            cells, any_value = [], False
            for cell in row:
                pos = (cell.row, cell.column)
                if pos in skip:
                    continue
                if cell.value is not None:
                    any_value = True
                span = span_at.get(pos)
                attrs = ""
                if span:
                    rs, cs = span
                    attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
                text = "" if cell.value is None else str(cell.value)
                cells.append(f"<td{attrs}>{escape(text)}</td>")
            if any_value:
                rows.append("<tr>" + "".join(cells) + "</tr>")
                row_count += 1
        out.append('<table class="ts-doc-table">' + "".join(rows) + "</table>" if rows else "<p class='ts-sub'>Лист пуст.</p>")
    wb.close()
    return "\n".join(out)


def _pdf_html(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        text = (page.extract_text() or "").strip()
        if text:
            parts.append("<p>" + "<br>".join(escape(line) for line in text.splitlines() if line.strip()) + "</p>")
    if len(reader.pages) > MAX_PDF_PAGES:
        parts.append("<p>…</p>")
    return "\n".join(parts) or "<p class='ts-sub'>В PDF нет извлекаемого текста (возможно, скан).</p>"


def extract_zip_entry(archive_data: bytes, path: str) -> bytes | None:
    """Достаём один файл из архива по имени (провал внутрь многофайлового zip)."""
    try:
        with zipfile.ZipFile(io.BytesIO(archive_data)) as zf:
            return zf.read(path)
    except (zipfile.BadZipFile, KeyError):
        return None


def extract_preview(data: bytes, filename: str) -> dict:
    """Возвращает {'kind': ..., 'html': ...} или {'kind': ..., 'error': ...}."""
    name = filename or ""
    is_zip = data[:2] == b"PK"

    # .zip-обёртка: достаём единственный вложенный файл
    if is_zip and (name.lower().endswith(".zip") or _ext(name) not in ("docx", "xlsx")):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = [n for n in zf.namelist() if not n.endswith("/")]
                office = [n for n in names if n.lower().rsplit(".", 1)[-1] in ("docx", "xlsx", "doc", "pdf")]
                if len(office) == 1:
                    return extract_preview(zf.read(office[0]), office[0])
                if names:
                    clickable = {n for n in office}
                    clickable |= {n for n in names if n.lower().endswith(".zip")}
                    listing = "".join(f"<li>{escape(n)}</li>" for n in names[:50])
                    return {"kind": "zip", "html": f"<p>Архив, файлы внутри:</p><ul>{listing}</ul>",
                            "zip_entries": [n for n in names if n in clickable]}
        except zipfile.BadZipFile:
            pass

    ext = _ext(name)
    try:
        if ext == "docx" or (is_zip and b"word/document.xml" in data[:4000] + data[-4000:]):
            return {"kind": "docx", "html": _docx_html(data)}
        if ext == "xlsx":
            return {"kind": "xlsx", "html": _xlsx_html(data)}
        if ext == "pdf" or data[:4] == b"%PDF":
            return {"kind": "pdf", "html": _pdf_html(data)}
        if ext == "doc":
            return {"kind": "doc", "error": "Старый формат .doc — предпросмотр недоступен, скачайте файл."}
    except Exception as exc:  # noqa: BLE001 — любой сбой парсинга -> мягкая ошибка
        return {"kind": ext or "?", "error": f"Не удалось разобрать файл: {exc}"}
    return {"kind": ext or "?", "error": "Формат не поддерживается для предпросмотра."}
