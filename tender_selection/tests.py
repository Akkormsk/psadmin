from unittest import mock

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from . import gosplan
from .filtering import match_title, parse_terms
from .models import FilterSettings, FoundTender, Organization, PullRun
from .services import run_pull


class CleanTitleTests(TestCase):
    def test_strips_eis_prefix(self):
        self.assertEqual(
            gosplan.clean_title("2026-06422**Техническое обслуживание медтехники"),
            "Техническое обслуживание медтехники",
        )

    def test_strips_single_asterisk_prefix(self):
        self.assertEqual(
            gosplan.clean_title("2026-06393*Поставка лекарственного препарата"),
            "Поставка лекарственного препарата",
        )

    def test_leaves_plain_title_untouched(self):
        self.assertEqual(gosplan.clean_title("Поставка кружек"), "Поставка кружек")


class EisUrlTests(TestCase):
    def test_known_type_maps_to_notice_segment(self):
        url = gosplan.eis_url("0131200001026005999", "epNotificationEF2020")
        self.assertIn("/notice/ea44/view/common-info.html?regNumber=0131200001026005999", url)

    def test_unknown_type_falls_back_to_search(self):
        url = gosplan.eis_url("123", "somethingElse")
        self.assertIn("extendedsearch/results.html?searchString=123", url)


class RunPullTests(TestCase):
    RECORD = {
        "purchase_number": "0131200001026005999",
        "object_info": "2026-06422**Поставка сувенирной продукции",
        "max_price": "377156.06",
        "currency_code": "RUB",
        "customers": ["3666018060"],
        "okpd2": ["32.99.12"],
        "region": 36,
        "stage": 1,
        "purchase_type": "epNotificationEF2020",
        "published_at": "2026-09-07T14:59:35.309000",
        "collecting_finished_at": "2026-09-16T06:00:00",
    }

    def test_creates_then_updates_and_logs(self):
        with mock.patch.object(gosplan, "iter_purchases", return_value=iter([self.RECORD])), \
             mock.patch("tender_selection.services.enrich_organizations", return_value=0):
            run = run_pull(days=3, max_requests=1)
        self.assertTrue(run.ok)
        self.assertEqual(run.created_count, 1)
        tender = FoundTender.objects.get(purchase_number="0131200001026005999")
        self.assertEqual(tender.title, "Поставка сувенирной продукции")
        self.assertEqual(tender.customer_inn, "3666018060")
        self.assertIn("ea44", tender.eis_url)

        with mock.patch.object(gosplan, "iter_purchases", return_value=iter([self.RECORD])), \
             mock.patch("tender_selection.services.enrich_organizations", return_value=0):
            run2 = run_pull(days=3, max_requests=1)
        self.assertEqual(run2.created_count, 0)
        self.assertEqual(run2.updated_count, 1)
        self.assertEqual(FoundTender.objects.count(), 1)

    def test_api_error_is_recorded_not_raised(self):
        with mock.patch.object(gosplan, "iter_purchases", side_effect=gosplan.GosplanError("HTTP 429")):
            run = run_pull(days=3, max_requests=1)
        self.assertFalse(run.ok)
        self.assertIn("429", run.error)

    def test_pull_never_triggers_risk_assessment_automatically(self):
        """Оценка рисков — платный запрос к ИИ-шлюзу за документами закупки, поэтому
        запускается только явной кнопкой на карточке (см. views.risk_status), а не
        сама по себе при каждой синхронизации новых тендеров."""
        settings = FilterSettings.load()
        settings.include_words = "сувенир"
        settings.save(update_fields=["include_words"])
        record = {**self.RECORD, "collecting_finished_at": None}
        with mock.patch.object(gosplan, "iter_purchases", side_effect=[iter([record]), iter([record])]), \
             mock.patch("tender_selection.services.notification_for", return_value={"source": {}}) as notification, \
             mock.patch("tender_selection.services.risk_assessment_for") as assess:
            run_pull(days=3, max_requests=1, classifiers=["32.99"])
            run_pull(days=3, max_requests=1, classifiers=["32.99"])
        notification.assert_not_called()
        assess.assert_not_called()

    def test_new_tender_outside_saved_filters_skips_risk_assessment(self):
        settings = FilterSettings.load()
        settings.include_words = "полиграфия"
        settings.save(update_fields=["include_words"])
        with mock.patch.object(gosplan, "iter_purchases", return_value=iter([self.RECORD])), \
             mock.patch("tender_selection.services.notification_for") as notification, \
             mock.patch("tender_selection.services.risk_assessment_for") as assess:
            run_pull(days=3, max_requests=1)
        notification.assert_not_called()
        assess.assert_not_called()

    def test_low_price_expired_and_excluded_tenders_skip_risk_assessment(self):
        settings = FilterSettings.load()
        settings.include_words = "сувенир"
        settings.exclude_words = "пластик"
        settings.save(update_fields=["include_words", "exclude_words"])
        records = [
            {**self.RECORD, "purchase_number": "1", "max_price": "100000", "collecting_finished_at": None},
            {**self.RECORD, "purchase_number": "2", "collecting_finished_at": "2020-01-01T00:00:00"},
            {**self.RECORD, "purchase_number": "3", "object_info": "Сувениры из пластика", "collecting_finished_at": None},
        ]
        with mock.patch.object(gosplan, "iter_purchases", return_value=iter(records)), \
             mock.patch("tender_selection.services.notification_for") as notification, \
             mock.patch("tender_selection.services.risk_assessment_for") as assess:
            run = run_pull(days=3, max_requests=1, classifiers=["32.99"])
        self.assertEqual(run.created_count, 3)
        notification.assert_not_called()
        assess.assert_not_called()


class RetryPendingRisksTests(TestCase):
    def test_retries_recent_matching_tender_after_notification_failure(self):
        from datetime import timedelta

        from .services import retry_pending_risks

        settings = FilterSettings.load()
        settings.include_words = "сувенир"
        settings.save(update_fields=["include_words"])
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="Сувенирная продукция", law="fz44",
            max_price=400000, last_pulled_at=timezone.now(), review=FoundTender.INTERESTING,
            notification_checked_at=timezone.now() - timedelta(hours=1),
        )
        FoundTender.objects.create(
            purchase_number="2", object_info="Медикаменты", law="fz44",
            max_price=400000, last_pulled_at=timezone.now(), review=FoundTender.INTERESTING,
        )
        with mock.patch("tender_selection.services.notification_for", return_value=NOTIFICATION_FIXTURE) as notification, \
             mock.patch("tender_selection.services.risk_assessment_for", return_value={"legal_risks": "ok"}) as assess:
            attempted, succeeded = retry_pending_risks()
        self.assertEqual((attempted, succeeded), (1, 1))
        notification.assert_called_once_with(tender, force=True)
        assess.assert_called_once_with(tender)

    def test_skips_unreviewed_tender(self):
        """На «Входящих» (review не тронут) риск ещё не актуален — не считаем,
        пока тендер не отправят «На оценку рисков»."""
        from .services import retry_pending_risks

        settings = FilterSettings.load()
        settings.include_words = "сувенир"
        settings.save(update_fields=["include_words"])
        FoundTender.objects.create(
            purchase_number="1", object_info="Сувенирная продукция", law="fz44",
            max_price=400000, last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.risk_assessment_for") as assess:
            attempted, succeeded = retry_pending_risks()
        self.assertEqual((attempted, succeeded), (0, 0))
        assess.assert_not_called()


class MultiSourceTests(TestCase):
    def test_fz223_record_shape(self):
        from tender_selection.services import _record_to_fields
        rec = {
            "purchase_number": "31900000001", "object_info": "Поставка сувенирной продукции",
            "customer": "7700000000", "submission_close_at": "2026-09-20T09:00:00",
            "max_price": "500000", "region": 77, "okpd2": ["32.99.11"],
        }
        f = _record_to_fields(rec, timezone.now(), "fz223")
        self.assertEqual(f["customer_inn"], "7700000000")
        self.assertIsNotNone(f["collecting_finished_at"])
        self.assertIn("223/purchase", f["eis_url"])

    def test_223_org_normalization(self):
        from tender_selection.services import _org_fields
        src = {"mainInfo": {"fullName": 'ПАО "РОССЕТИ ЦЕНТР"', "shortName": 'ПАО "РОССЕТИ ЦЕНТР"',
                            "postalAddress": "г. Москва, ул. Ордынка", "region": "Г.МОСКВА"},
               "contactInfo": {"contactEmail": "x@mrsk.ru", "website": "mrsk-1.ru"}}
        f = _org_fields(src, "fz223")
        self.assertIn("РОССЕТИ ЦЕНТР", f["name"])
        self.assertEqual(f["email"], "x@mrsk.ru")
        self.assertEqual(f["region_name"], "Г.Москва")

    def test_list_law_filter(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        FilterSettings.objects.update_or_create(pk=1, defaults={"min_price": 0})
        FoundTender.objects.create(purchase_number="1", law="fz44", object_info="A", title="Кружка 44", last_pulled_at=timezone.now())
        FoundTender.objects.create(purchase_number="2", law="fz223", object_info="B", title="Кружка 223", last_pulled_at=timezone.now())
        r44 = self.client.get(reverse("tender_selection:list") + "?view=list&law=fz44")
        self.assertContains(r44, "Кружка 44")
        self.assertNotContains(r44, "Кружка 223")
        rboth = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertContains(rboth, "Кружка 44")
        self.assertContains(rboth, "Кружка 223")

    def test_settings_save_categories_regions_laws(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        self.client.post(reverse("tender_selection:settings"), {
            "include_words": "", "exclude_words": "", "min_price": "300000", "window_days": "7",
            "okpd2": ["18.1", "32.99"], "region": ["77", "78"], "law": ["fz44", "fz223"],
        })
        s = FilterSettings.load()
        self.assertEqual(sorted(s.okpd2_codes), ["18.1", "32.99"])
        self.assertEqual(sorted(s.regions), ["77", "78"])
        self.assertEqual(sorted(s.laws), ["fz223", "fz44"])


class PushToEstimateTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.client.force_login(self.admin)
        for name in ("fetch_clarifications", "fetch_complaints"):
            patcher = mock.patch.object(gosplan, name, return_value=[])
            patcher.start()
            self.addCleanup(patcher.stop)
        self.tender = FoundTender.objects.create(
            purchase_number="1", law="fz44", object_info="x", title="Сувенирка",
            max_price=500000, last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )

    def test_creates_estimate_with_lines_from_notification(self):
        resp = self.client.post(reverse("tender_selection:push", args=[self.tender.pk]))
        from tenders.models import TenderEstimate
        est = TenderEstimate.objects.get()
        self.assertEqual(est.tender_number, "1")
        self.assertIn("ПРИМЕР", est.name)
        self.assertNotIn("№", est.name)
        self.assertEqual(est.lines.count(), 1)
        line = est.lines.first()
        self.assertEqual(line.name, "Баннер")
        self.assertEqual(str(line.nmck_unit), "3700.00")
        reqs = line.requirements["requirements"]
        self.assertTrue(any("баннерная ткань" in r["value"] for r in reqs))
        self.tender.refresh_from_db()
        self.assertEqual(self.tender.status, FoundTender.PUSHED)
        self.assertEqual(self.tender.pushed_estimate_id, est.pk)
        self.assertRedirects(resp, f"/tenders/{est.pk}/", fetch_redirect_response=False)

    def test_pushed_tender_moves_to_calculation_column_on_kanban(self):
        """По умолчанию tender_selection:list — канбан: запушенный тендер
        пропадает из Входящие/Проверка (FoundTender.status=PUSHED больше не
        попадает в этот запрос) и появляется карточкой просчёта в «Расчёте»."""
        self.client.post(reverse("tender_selection:push", args=[self.tender.pk]))
        from tenders.models import TenderEstimate

        est = TenderEstimate.objects.get()
        kanban_resp = self.client.get(reverse("tender_selection:list"))
        # "1" (purchase_number) сам по себе слишком общая строка (встречается
        # в разметке независимо от карточек) — проверяем заголовок тендера,
        # которого не должно остаться среди карточек «Входящие»/«Проверка».
        self.assertNotContains(kanban_resp, self.tender.title)
        self.assertContains(kanban_resp, f"№ {est.tender_number}")
        detail_resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertContains(detail_resp, "На расчёте — открыть просчёт")
        self.assertNotContains(detail_resp, 'name="review"')  # селектор статуса скрыт

    def test_pushed_tender_still_shows_pill_in_flat_list_view(self):
        """Старое поведение списка (?view=list) никуда не делось, просто больше
        не значение по умолчанию — тумблер Канбан/Список должен продолжать работать."""
        self.client.post(reverse("tender_selection:push", args=[self.tender.pk]))
        list_resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertContains(list_resp, "ts-onestimate-pill")
        self.assertContains(list_resp, "На расчёте")

    def test_second_push_opens_existing(self):
        self.client.post(reverse("tender_selection:push", args=[self.tender.pk]))
        from tenders.models import TenderEstimate
        first = TenderEstimate.objects.get().pk
        self.client.post(reverse("tender_selection:push", args=[self.tender.pk]))
        self.assertEqual(TenderEstimate.objects.count(), 1)


class DocumentPreviewTests(TestCase):
    def _docx_bytes(self):
        import io
        from docx import Document
        d = Document()
        d.add_paragraph("Описание объекта закупки")
        t = d.add_table(rows=2, cols=2)
        t.rows[0].cells[0].text = "Параметр"
        t.rows[0].cells[1].text = "Значение"
        t.rows[1].cells[0].text = "Материал"
        t.rows[1].cells[1].text = "хлопок 100%"
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def _xlsx_bytes(self):
        import io
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Позиция", "Кол-во", "Цена"])
        ws.append(["Кружка", 100, 250])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def test_extract_docx(self):
        from .documents import extract_preview
        r = extract_preview(self._docx_bytes(), "ООЗ.docx")
        self.assertEqual(r["kind"], "docx")
        self.assertIn("Описание объекта закупки", r["html"])
        self.assertIn("хлопок 100%", r["html"])
        self.assertIn("ts-doc-table", r["html"])

    def test_extract_xlsx(self):
        from .documents import extract_preview
        r = extract_preview(self._xlsx_bytes(), "НМЦК.xlsx")
        self.assertEqual(r["kind"], "xlsx")
        self.assertIn("Кружка", r["html"])
        self.assertIn("<table", r["html"])

    def test_extract_docx_table_with_merged_cells(self):
        from docx import Document
        import io

        from .documents import extract_preview

        d = Document()
        t = d.add_table(rows=3, cols=3)
        t.cell(0, 0).merge(t.cell(0, 2))
        t.cell(0, 0).text = "Спецификация"
        t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = "Наименование", "Кол-во", "Цена"
        t.cell(1, 0).merge(t.cell(2, 0))
        t.cell(1, 0).text = "Кружка"
        t.cell(2, 1).text, t.cell(2, 2).text = "100", "250"
        buf = io.BytesIO()
        d.save(buf)

        r = extract_preview(buf.getvalue(), "Смета.docx")
        self.assertEqual(r["kind"], "docx")
        self.assertIn('colspan="3"', r["html"])  # заголовок на всю ширину
        self.assertIn('rowspan="2"', r["html"])  # "Кружка" на две строки
        self.assertIn("Спецификация", r["html"])
        # "Кружка" — объединённая по вертикали ячейка, должна попасть в HTML РОВНО один раз
        # (регрессия: python-docx схлопывает row.cells для vMerge-продолжений в тот же
        # объект, что и ячейка-шапка, из-за чего текст дублировался на каждой строке).
        self.assertEqual(r["html"].count("Кружка"), 1)
        # строка-продолжение объединения (третья строка таблицы, "100"/"250") должна
        # содержать только 2 <td>, а не 3 — если бы rowspan не сработал, тут была бы
        # лишняя ячейка с "Кружка".
        continuation_row = r["html"].split("<tr>")[3]
        self.assertEqual(continuation_row.count("<td"), 2)

    def test_extract_xlsx_with_merged_cells(self):
        from openpyxl import Workbook
        import io

        from .documents import extract_preview

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Смета"
        ws.merge_cells("A1:C1")
        ws.append(["Наименование", "Кол-во", "Цена"])
        ws["A3"] = "Кружка"
        ws.merge_cells("A3:A4")
        ws["B3"], ws["C3"] = 60, 250
        ws["B4"], ws["C4"] = 40, 250
        buf = io.BytesIO()
        wb.save(buf)

        r = extract_preview(buf.getvalue(), "Смета.xlsx")
        self.assertEqual(r["kind"], "xlsx")
        self.assertIn('colspan="3"', r["html"])  # заголовок на всю ширину
        self.assertIn('rowspan="2"', r["html"])  # «Кружка» на две строки
        self.assertIn("Смета", r["html"])

    def test_extract_rejects_unknown(self):
        from .documents import extract_preview
        r = extract_preview(b"random bytes", "notes.txt")
        self.assertIn("error", r)

    def _zip_with(self, files: dict) -> bytes:
        import io
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for name, data in files.items():
                zf.writestr(name, data)
        return buf.getvalue()

    def test_extract_zip_with_multiple_files_lists_clickable_entries(self):
        from .documents import extract_preview
        archive = self._zip_with({
            "Пр1.docx": self._docx_bytes(),
            "Пр2.xlsx": self._xlsx_bytes(),
            "картинка.jpg": b"not-really-an-image",
        })
        r = extract_preview(archive, "bundle.zip")
        self.assertEqual(r["kind"], "zip")
        self.assertIn("Пр1.docx", r["zip_entries"])
        self.assertIn("Пр2.xlsx", r["zip_entries"])
        self.assertNotIn("картинка.jpg", r["zip_entries"])

    def test_extract_zip_entry_returns_bytes(self):
        from .documents import extract_zip_entry
        archive = self._zip_with({"inner.docx": self._docx_bytes()})
        data = extract_zip_entry(archive, "inner.docx")
        self.assertEqual(data, self._docx_bytes())

    def test_extract_zip_entry_missing_returns_none(self):
        from .documents import extract_zip_entry
        archive = self._zip_with({"inner.docx": self._docx_bytes()})
        self.assertIsNone(extract_zip_entry(archive, "absent.docx"))

    def test_extract_zip_entry_bad_archive_returns_none(self):
        from .documents import extract_zip_entry
        self.assertIsNone(extract_zip_entry(b"not a zip", "inner.docx"))

    def test_fetch_document_blocks_foreign_url(self):
        from .documents import DocumentError, fetch_document
        with self.assertRaises(DocumentError):
            fetch_document("https://evil.example/x.docx")

    def test_view_fetches_extracts_and_caches(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", return_value=self._docx_bytes()) as f:
            resp = self.client.get(reverse("tender_selection:doc_preview", args=[tender.pk, 0]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Описание объекта закупки", resp.json()["html"])
        f.assert_called_once()
        # second call served from cache — no fetch
        with mock.patch("tender_selection.services.fetch_document_via_eis") as f2:
            self.client.get(reverse("tender_selection:doc_preview", args=[tender.pk, 0]))
        f2.assert_not_called()

    def test_view_reports_fetch_error(self):
        from .documents import DocumentError
        from .eis_docs import EisDocsError
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", side_effect=EisDocsError("токен не задан")), \
             mock.patch("tender_selection.services.fetch_document", side_effect=DocumentError("ЕИС недоступен")):
            resp = self.client.get(reverse("tender_selection:doc_preview", args=[tender.pk, 0]))
        self.assertIn("токен не задан", resp.json()["error"])
        self.assertIn("ЕИС недоступен", resp.json()["error"])
        from .models import DocumentPreview
        self.assertFalse(DocumentPreview.objects.exists())  # сетевой сбой не кэшируется

    def test_view_falls_back_to_direct_link_when_eis_fails(self):
        from .eis_docs import EisDocsError
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="123", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", side_effect=EisDocsError("лимит ЕИС исчерпан")) as via_eis, \
             mock.patch("tender_selection.services.fetch_document", return_value=self._docx_bytes()) as direct:
            resp = self.client.get(reverse("tender_selection:doc_preview", args=[tender.pk, 0]))
        via_eis.assert_called_once_with("123", mock.ANY)
        direct.assert_called_once()
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Описание объекта закупки", resp.json()["html"])
        from .models import DocumentPreview
        self.assertTrue(DocumentPreview.objects.exists())  # успешный резервный путь кэшируется

    def test_view_uses_eis_without_touching_direct_link_when_it_works(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", return_value=self._docx_bytes()), \
             mock.patch("tender_selection.services.fetch_document") as direct:
            resp = self.client.get(reverse("tender_selection:doc_preview", args=[tender.pk, 0]))
        direct.assert_not_called()
        self.assertIn("Описание объекта закупки", resp.json()["html"])

    def test_doc_zip_entry_returns_inner_file(self):
        archive = self._zip_with({"Пр1.docx": self._docx_bytes(), "Пр2.xlsx": self._xlsx_bytes()})
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", return_value=archive):
            resp = self.client.get(reverse("tender_selection:doc_zip_entry", args=[tender.pk, 0, "Пр1.docx"]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Описание объекта закупки", resp.json()["html"])

    def test_doc_zip_entry_missing_file_reports_error(self):
        archive = self._zip_with({"Пр1.docx": self._docx_bytes()})
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.services.fetch_document_via_eis", return_value=archive):
            resp = self.client.get(reverse("tender_selection:doc_zip_entry", args=[tender.pk, 0, "absent.docx"]))
        self.assertIn("не найден", resp.json()["error"])

    def test_eis_diag_reports_probe_results(self):
        import socket as socket_mod

        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))

        def fake_connect(addr, timeout=None):
            host = addr[0]
            if "gosplan" in host:
                return mock.Mock(close=lambda: None)
            raise TimeoutError("timed out")

        with mock.patch.object(socket_mod, "create_connection", side_effect=fake_connect), \
             mock.patch("urllib.request.urlopen", side_effect=TimeoutError("timed out")):
            resp = self.client.get(reverse("tender_selection:eis_diag"))
        self.assertEqual(resp.status_code, 200)
        results = resp.json()["results"]
        self.assertEqual(len(results), 9)
        by_probe = {r["probe"]: r for r in results}
        self.assertTrue(any("gosplan" in k and v["ok"] for k, v in by_probe.items()))
        self.assertTrue(any("zakupki" in k and not v["ok"] for k, v in by_probe.items()))

    def test_eis_diag_requires_superuser(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_user("u", "u@e.ru", "p"))
        resp = self.client.get(reverse("tender_selection:eis_diag"))
        self.assertEqual(resp.status_code, 403)

    def test_doc_upload_parses_and_caches(self):
        from django.core.files.uploadedfile import SimpleUploadedFile

        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        upload = SimpleUploadedFile("ООЗ.docx", self._docx_bytes())
        resp = self.client.post(reverse("tender_selection:doc_upload", args=[tender.pk, 0]), {"file": upload})
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Описание объекта закупки", resp.json()["html"])
        from .models import DocumentPreview
        self.assertTrue(DocumentPreview.objects.exists())

    def test_doc_upload_requires_file(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        resp = self.client.post(reverse("tender_selection:doc_upload", args=[tender.pk, 0]))
        self.assertIn("Файл не выбран", resp.json()["error"])

    def test_doc_upload_requires_superuser_or_tender_owner(self):
        # 404, не 403 — доступ теперь per-tender (см. TenderViewerAccessTests), а
        # тендера pk=1 в этой изолированной БД теста вообще нет.
        User = get_user_model()
        self.client.force_login(User.objects.create_user("u", "u@e.ru", "p"))
        resp = self.client.post(reverse("tender_selection:doc_upload", args=[1, 0]))
        self.assertEqual(resp.status_code, 404)

    def test_doc_upload_rejects_get(self):
        User = get_user_model()
        self.client.force_login(User.objects.create_superuser("a", "a@e.ru", "p"))
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", last_pulled_at=timezone.now(),
            notification_raw=NOTIFICATION_FIXTURE,
        )
        resp = self.client.get(reverse("tender_selection:doc_upload", args=[tender.pk, 0]))
        self.assertEqual(resp.status_code, 405)


def _build_minimal_doc(word_document_stream: bytes) -> bytes:
    """Собирает МИНИМАЛЬНЫЙ валидный OLE/CFB-контейнер с одним потоком
    "WordDocument" — без Word/LibreOffice под рукой это единственный способ
    честно проверить разбор .doc на настоящей структуре контейнера (round-trip
    проверен через olefile при разработке). Внутренний формат FIB Word не
    эмулируем — экстрактор его не разбирает, только сканирует сырые байты потока."""
    import struct

    SECTOR = 512
    ENDOFCHAIN, FREESECT, FATSECT, NOSTREAM = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF

    def name_field(name):
        raw = (name + "\0").encode("utf-16-le")
        return raw + b"\0" * (64 - len(raw))

    size = max(4096, ((len(word_document_stream) + SECTOR - 1) // SECTOR) * SECTOR)
    stream = word_document_stream + b"\0" * (size - len(word_document_stream))
    n_stream_sectors = size // SECTOR
    fat_idx, dir_idx, first_stream = 0, 1, 2
    total = 2 + n_stream_sectors

    fat = [FREESECT] * 128
    fat[fat_idx], fat[dir_idx] = FATSECT, ENDOFCHAIN
    for i in range(n_stream_sectors):
        sec = first_stream + i
        fat[sec] = (sec + 1) if i < n_stream_sectors - 1 else ENDOFCHAIN
    fat_bytes = b"".join(struct.pack("<I", v) for v in fat)

    def dir_entry(name, obj_type, child, start_sector, stream_size):
        raw_name = (name + "\0").encode("utf-16-le") if name else b""
        return (
            name_field(name) + struct.pack("<H", len(raw_name)) + struct.pack("<B", obj_type) +
            struct.pack("<B", 1) + struct.pack("<I", NOSTREAM) + struct.pack("<I", NOSTREAM) +
            struct.pack("<I", child) + b"\0" * 16 + struct.pack("<I", 0) + struct.pack("<Q", 0) +
            struct.pack("<Q", 0) + struct.pack("<I", start_sector) + struct.pack("<Q", stream_size)
        )

    directory = (
        dir_entry("Root Entry", 5, 1, ENDOFCHAIN, 0) +
        dir_entry("WordDocument", 2, NOSTREAM, first_stream, len(stream)) +
        dir_entry("", 0, NOSTREAM, 0, 0) + dir_entry("", 0, NOSTREAM, 0, 0)
    )

    difat = [fat_idx] + [FREESECT] * 108
    header = (
        bytes.fromhex("d0cf11e0a1b11ae1") + b"\0" * 16 +
        struct.pack("<H", 0x003E) + struct.pack("<H", 0x0003) + struct.pack("<H", 0xFFFE) +
        struct.pack("<H", 9) + struct.pack("<H", 6) + b"\0" * 6 +
        struct.pack("<I", 0) + struct.pack("<I", 1) + struct.pack("<I", dir_idx) +
        struct.pack("<I", 0) + struct.pack("<I", 0x1000) + struct.pack("<I", ENDOFCHAIN) +
        struct.pack("<I", 0) + struct.pack("<I", ENDOFCHAIN) + struct.pack("<I", 0) +
        b"".join(struct.pack("<I", v) for v in difat)
    )

    body = bytearray(b"\0" * (total * SECTOR))
    body[fat_idx * SECTOR:(fat_idx + 1) * SECTOR] = fat_bytes
    body[dir_idx * SECTOR:(dir_idx + 1) * SECTOR] = directory
    body[first_stream * SECTOR:(first_stream + n_stream_sectors) * SECTOR] = stream
    return header + bytes(body)


class LegacyDocTests(TestCase):
    """Старый бинарный .doc (OLE) — без Word/LibreOffice, но на настоящей структуре
    контейнера (см. _build_minimal_doc). Оба возможных варианта хранения текста
    в WordDocument, плюс отказоустойчивость на мусоре."""

    def test_extracts_utf16_text(self):
        from .documents import extract_preview

        text = "Спецификация поставки\rКружка сувенирная — 100 шт.\r"
        data = _build_minimal_doc(text.encode("utf-16-le"))
        r = extract_preview(data, "Извещение.doc")
        self.assertEqual(r["kind"], "doc")
        self.assertIn("Спецификация поставки", r["html"])
        self.assertIn("Кружка сувенирная", r["html"])

    def test_extracts_cp1251_text(self):
        """Word 97 часто сохранял русский текст однобайтовым CP1251, а не UTF-16 —
        экстрактор обязан опознать и этот вариант, не только Unicode."""
        from .documents import extract_preview

        text = "Техническое задание\rПоставка канцелярских товаров\r"
        data = _build_minimal_doc(text.encode("cp1251"))
        r = extract_preview(data, "ТЗ.doc")
        self.assertEqual(r["kind"], "doc")
        self.assertIn("Техническое задание", r["html"])
        self.assertIn("Поставка канцелярских товаров", r["html"])

    def test_non_ole_garbage_is_soft_error_not_crash(self):
        from .documents import extract_preview

        r = extract_preview(b"just some random short bytes", "fake.doc")
        self.assertEqual(r["kind"], "doc")
        self.assertIn("error", r)

    def test_empty_stream_reports_no_text_instead_of_crashing(self):
        from .documents import extract_preview

        r = extract_preview(_build_minimal_doc(b""), "empty.doc")
        self.assertEqual(r["kind"], "doc")
        self.assertIn("нет извлекаемого текста", r["html"])


class EisDocsTests(TestCase):
    """Официальный резервный канал ЕИС (getDocsIP) — без сети, всё замокано."""

    def _zip_with(self, files: dict) -> bytes:
        import io
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            for name, data in files.items():
                zf.writestr(name, data)
        return buf.getvalue()

    def test_fetch_archive_urls_requires_token(self):
        from . import eis_docs
        with mock.patch.dict("os.environ", {"EIS_TOKEN": ""}):
            with self.assertRaises(eis_docs.EisDocsError):
                eis_docs.fetch_archive_urls("123")

    def test_fetch_archive_urls_parses_response(self):
        from . import eis_docs
        soap_response = (
            "<soapenv:Envelope xmlns:soapenv='x'><soapenv:Body>"
            "<ns2:getDocsByReestrNumberResponse><dataInfo>"
            "<archiveUrl>https://int44.zakupki.gov.ru/archive/1.zip</archiveUrl>"
            "</dataInfo></ns2:getDocsByReestrNumberResponse>"
            "</soapenv:Body></soapenv:Envelope>"
        )
        with mock.patch.dict("os.environ", {"EIS_TOKEN": "tok"}), \
             mock.patch.object(eis_docs, "_soap_call", return_value=soap_response):
            urls = eis_docs.fetch_archive_urls("123")
        self.assertEqual(urls, ["https://int44.zakupki.gov.ru/archive/1.zip"])

    def test_fetch_archive_urls_raises_on_fault(self):
        from . import eis_docs
        fault = "<soapenv:Fault><faultstring>Неверный токен</faultstring></soapenv:Fault>"
        with mock.patch.dict("os.environ", {"EIS_TOKEN": "tok"}), \
             mock.patch.object(eis_docs, "_soap_call", return_value=fault):
            with self.assertRaises(eis_docs.EisDocsError) as ctx:
                eis_docs.fetch_archive_urls("123")
        self.assertIn("Неверный токен", str(ctx.exception))

    def test_find_file_matches_by_normalized_name(self):
        from . import eis_docs
        archive = self._zip_with({"docs/Описание объекта закупки.docx": b"content-a", "other.pdf": b"x"})
        found = eis_docs.find_file(archive, "  Описание   объекта закупки.docx")
        self.assertIsNotNone(found)
        self.assertEqual(found[0], b"content-a")

    def test_find_file_looks_inside_nested_zip(self):
        from . import eis_docs
        inner = self._zip_with({"target.pdf": b"inner-bytes"})
        outer = self._zip_with({"bundle.zip": inner})
        found = eis_docs.find_file(outer, "target.pdf")
        self.assertEqual(found[0], b"inner-bytes")

    def test_find_file_returns_none_when_absent(self):
        from . import eis_docs
        archive = self._zip_with({"other.pdf": b"x"})
        self.assertIsNone(eis_docs.find_file(archive, "missing.docx"))

    def test_fetch_document_via_eis_tries_next_archive_on_failure(self):
        from . import eis_docs
        with mock.patch.object(eis_docs, "fetch_archive_urls", return_value=["u1", "u2"]), \
             mock.patch.object(eis_docs, "download_archive", side_effect=[
                 eis_docs.EisDocsError("первый архив недоступен"),
                 self._zip_with({"file.docx": b"ok"}),
             ]):
            data = eis_docs.fetch_document_via_eis("123", "file.docx")
        self.assertEqual(data, b"ok")

    def test_fetch_document_via_eis_raises_when_not_found_anywhere(self):
        from . import eis_docs
        with mock.patch.object(eis_docs, "fetch_archive_urls", return_value=["u1"]), \
             mock.patch.object(eis_docs, "download_archive", return_value=self._zip_with({"other.pdf": b"x"})):
            with self.assertRaises(eis_docs.EisDocsError):
                eis_docs.fetch_document_via_eis("123", "missing.docx")

    def test_archive_loop_budget_stops_before_trying_every_archive(self):
        """Предохранитель от WORKER TIMEOUT (см. ARCHIVE_LOOP_BUDGET_SECONDS): если
        перебор архивов не укладывается в бюджет, дальнейшие архивы не трогаем — иначе
        при 7+ архивах по 90с каждый легко перевешиваем gunicorn --timeout 600 на
        единственном sync-воркере (см. серию WORKER TIMEOUT 14-15.09.2026)."""
        from . import eis_docs
        import time as time_module

        def slow_fail(url):
            time_module.sleep(0.05)
            raise eis_docs.EisDocsError("не вышло")

        with mock.patch.object(eis_docs, "ARCHIVE_LOOP_BUDGET_SECONDS", 0.08), \
             mock.patch.object(eis_docs, "fetch_archive_urls", return_value=["u1", "u2", "u3", "u4", "u5"]), \
             mock.patch.object(eis_docs, "download_archive", side_effect=slow_fail) as dl:
            with self.assertRaises(eis_docs.EisDocsError) as ctx:
                eis_docs.fetch_document_via_eis("123", "file.docx")
        self.assertLess(dl.call_count, 5)  # не дошёл до всех архивов
        self.assertIn("долго", str(ctx.exception))

    def test_archive_loop_within_budget_still_tries_all_archives(self):
        """Регрессия: обычный (быстрый) перебор нескольких архивов не должен пострадать
        от нового предохранителя — бюджет с большим запасом на реальные случаи."""
        from . import eis_docs
        with mock.patch.object(eis_docs, "fetch_archive_urls", return_value=["u1", "u2"]), \
             mock.patch.object(eis_docs, "download_archive", side_effect=[
                 eis_docs.EisDocsError("первый архив недоступен"),
                 self._zip_with({"file.docx": b"ok"}),
             ]):
            data = eis_docs.fetch_document_via_eis("123", "file.docx")
        self.assertEqual(data, b"ok")


class RetryPendingDocumentsTests(TestCase):
    """Фоновые повторы скачивания — сеть до ЕИС нестабильна (плавающая блокировка),
    поэтому периодически подбираем то, что не скачалось раньше."""

    def _docx_bytes(self):
        import io
        from docx import Document
        d = Document()
        d.add_paragraph("текст")
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_skips_already_cached_documents(self):
        from .models import DocumentPreview
        from .services import retry_pending_documents

        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )
        DocumentPreview.objects.create(
            url="https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=A",
            filename="cached", kind="docx", html="<p>уже есть</p>",
        )
        with mock.patch("tender_selection.eis_docs.fetch_document_via_eis") as via_eis:
            attempted, succeeded = retry_pending_documents()
        via_eis.assert_not_called()
        self.assertEqual((attempted, succeeded), (0, 0))

    def test_fetches_and_caches_uncached_document_via_eis(self):
        from .models import DocumentPreview
        from .services import retry_pending_documents

        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.eis_docs.fetch_document_via_eis", return_value=self._docx_bytes()) as via_eis:
            attempted, succeeded = retry_pending_documents()
        via_eis.assert_called_once_with("1", mock.ANY)
        self.assertEqual((attempted, succeeded), (1, 1))
        self.assertTrue(DocumentPreview.objects.filter(
            url="https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=A",
        ).exists())

    def test_falls_back_to_direct_link_and_does_not_cache_on_total_failure(self):
        from .documents import DocumentError
        from .eis_docs import EisDocsError
        from .models import DocumentPreview
        from .services import retry_pending_documents

        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )
        with mock.patch("tender_selection.eis_docs.fetch_document_via_eis", side_effect=EisDocsError("нет сети")), \
             mock.patch("tender_selection.documents.fetch_document", side_effect=DocumentError("нет сети")):
            attempted, succeeded = retry_pending_documents()
        self.assertEqual((attempted, succeeded), (1, 0))
        self.assertFalse(DocumentPreview.objects.exists())

    def test_respects_limit(self):
        from .services import retry_pending_documents

        for i in range(3):
            FoundTender.objects.create(
                purchase_number=str(i), object_info="x", title="T", law="fz44",
                last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
            )
        with mock.patch("tender_selection.eis_docs.fetch_document_via_eis", return_value=self._docx_bytes()):
            attempted, succeeded = retry_pending_documents(limit=1)
        self.assertEqual((attempted, succeeded), (1, 1))


class NotificationForTests(TestCase):
    """notification_for() должен помечать попытку (notification_checked_at) даже при
    сбое — иначе бейдж «⚠ нет данных» в списке не отличить от «ещё не проверяли»."""

    def test_success_sets_raw_and_checked_at(self):
        from .services import notification_for

        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(),
        )
        with mock.patch("tender_selection.gosplan.fetch_notification", return_value=NOTIFICATION_FIXTURE):
            payload = notification_for(tender)
        tender.refresh_from_db()
        self.assertEqual(payload, NOTIFICATION_FIXTURE)
        self.assertTrue(tender.notification_raw)
        self.assertIsNotNone(tender.notification_checked_at)

    def test_failure_still_sets_checked_at(self):
        from . import gosplan
        from .services import notification_for

        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(),
        )
        with mock.patch("tender_selection.gosplan.fetch_notification", side_effect=gosplan.GosplanError("нет сети")):
            payload = notification_for(tender)
        tender.refresh_from_db()
        self.assertIsNone(payload)
        self.assertFalse(tender.notification_raw)
        self.assertIsNotNone(tender.notification_checked_at)  # попытка была — это и есть реальный сбой


class RetryPendingNotificationsTests(TestCase):
    """Фоновая догрузка извещений для свежих тендеров — без неё бейдж «нет данных»
    в списке горел бы на каждом только что выгруженном тендере (см. views.py)."""

    def test_fetches_never_checked_tenders(self):
        from .services import retry_pending_notifications

        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            last_pulled_at=timezone.now(),
        )  # notification_checked_at пуст — ни разу не пробовали
        with mock.patch("tender_selection.gosplan.fetch_notification", return_value=NOTIFICATION_FIXTURE) as fetch:
            attempted, succeeded = retry_pending_notifications()
        fetch.assert_called_once_with("1")
        self.assertEqual((attempted, succeeded), (1, 1))
        tender.refresh_from_db()
        self.assertTrue(tender.notification_raw)

    def test_skips_already_checked_tenders(self):
        from .services import retry_pending_notifications

        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            notification_checked_at=timezone.now(), last_pulled_at=timezone.now(),
        )  # уже проверяли (успешно или нет) — фон это трогать не должен
        with mock.patch("tender_selection.gosplan.fetch_notification") as fetch:
            attempted, succeeded = retry_pending_notifications()
        fetch.assert_not_called()
        self.assertEqual((attempted, succeeded), (0, 0))

    def test_respects_limit(self):
        from .services import retry_pending_notifications

        for i in range(3):
            FoundTender.objects.create(
                purchase_number=str(i), object_info="x", title="T", law="fz44",
                last_pulled_at=timezone.now(),
            )
        with mock.patch("tender_selection.gosplan.fetch_notification", return_value=NOTIFICATION_FIXTURE):
            attempted, succeeded = retry_pending_notifications(limit=1)
        self.assertEqual((attempted, succeeded), (1, 1))

    def test_fz223_never_attempted(self):
        """У 223-ФЗ нет извещения по конструкции источника — фон не должен даже
        пытаться (notification_for сам вернёт None по law, но раз мы отбираем по
        checked_at, 223-ФЗ тендер так и останется checked_at=None навсегда — это
        нормально: у него просто никогда не будет notification_missing=True, см.
        views.py, где условие уже требует law == 'fz44')."""
        from .services import retry_pending_notifications

        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz223",
            last_pulled_at=timezone.now(),
        )
        with mock.patch("tender_selection.gosplan.fetch_notification") as fetch:
            attempted, succeeded = retry_pending_notifications()
        fetch.assert_not_called()
        self.assertEqual((attempted, succeeded), (0, 0))


class RetryPendingOutcomesTests(TestCase):
    def setUp(self):
        from tenders.models import TenderEstimate

        User = get_user_model()
        owner = User.objects.create_user("owner", "o@e.ru", "pw")
        self.estimate = TenderEstimate.objects.create(
            owner=owner, tender_number="0342300000126000995", name="Тест",
            status=TenderEstimate.PENDING, summary_snapshot={"nmck_total": "515400"},
        )

    def test_skips_already_checked(self):
        from .services import retry_pending_outcomes

        self.estimate.outcome_checked_at = timezone.now()
        self.estimate.save(update_fields=["outcome_checked_at"])
        with mock.patch("tender_selection.gosplan.fetch_contracts") as fetch:
            attempted, succeeded = retry_pending_outcomes()
        fetch.assert_not_called()
        self.assertEqual((attempted, succeeded), (0, 0))

    def test_not_found_leaves_estimate_untouched(self):
        from .services import retry_pending_outcomes

        with mock.patch("tender_selection.gosplan.fetch_contracts", return_value=[]):
            attempted, succeeded = retry_pending_outcomes()
        self.assertEqual((attempted, succeeded), (1, 0))
        self.estimate.refresh_from_db()
        self.assertEqual(self.estimate.status, self.estimate.PENDING)
        self.assertIsNone(self.estimate.outcome_checked_at)

    def test_found_without_company_inn_records_price_but_not_status(self):
        from .services import retry_pending_outcomes

        row = {"price": 386550, "suppliers": ["526001302080"]}
        with mock.patch("tender_selection.gosplan.fetch_contracts", return_value=[row]), \
             mock.patch.dict("os.environ", {}, clear=False):
            import os
            os.environ.pop("COMPANY_INN", None)
            attempted, succeeded = retry_pending_outcomes()
        self.assertEqual((attempted, succeeded), (1, 1))
        self.estimate.refresh_from_db()
        self.assertEqual(self.estimate.status, self.estimate.PENDING)
        self.assertEqual(str(self.estimate.actual_price), "386550.00")
        self.assertIsNotNone(self.estimate.outcome_checked_at)

    def test_found_with_matching_company_inn_marks_won(self):
        from .services import retry_pending_outcomes

        row = {"price": 386550, "suppliers": ["526001302080"]}
        with mock.patch("tender_selection.gosplan.fetch_contracts", return_value=[row]), \
             mock.patch.dict("os.environ", {"COMPANY_INN": "526001302080"}):
            attempted, succeeded = retry_pending_outcomes()
        self.assertEqual((attempted, succeeded), (1, 1))
        self.estimate.refresh_from_db()
        self.assertEqual(self.estimate.status, self.estimate.WON)
        self.assertEqual(self.estimate.outcome_source, self.estimate.OUTCOME_AUTO)

        from .models import ContractStat
        self.assertTrue(ContractStat.objects.filter(purchase_number=self.estimate.tender_number, is_ours=True).exists())

    def test_respects_limit(self):
        from tenders.models import TenderEstimate

        from .services import retry_pending_outcomes

        for i in range(3):
            TenderEstimate.objects.create(
                owner=self.estimate.owner, tender_number=f"200000000000000000{i}",
                name="Тест", status=TenderEstimate.PENDING,
            )
        with mock.patch("tender_selection.gosplan.fetch_contracts", return_value=[]):
            attempted, succeeded = retry_pending_outcomes(limit=1)
        self.assertEqual((attempted, succeeded), (1, 0))


class RiskAssessmentSelectDocumentsTests(TestCase):
    def test_prioritizes_contract_over_description(self):
        from .risk_assessment import select_documents

        docs = [
            {"name": "Описание объекта закупки.docx", "kind": "Описание объекта закупки", "url": "a"},
            {"name": "Проект контракта.docx", "kind": "Проект контракта", "url": "b"},
        ]
        selected = select_documents(docs)
        self.assertEqual([d["name"] for d in selected], ["Проект контракта.docx", "Описание объекта закупки.docx"])

    def test_skips_images_and_price_justification(self):
        from .risk_assessment import select_documents

        docs = [
            {"name": "Обоснование НМЦК.docx", "kind": "Обоснование начальной (максимальной) цены контракта", "url": "a"},
            {"name": "Фото образца.jpg", "kind": "Изображение", "url": "b"},
        ]
        self.assertEqual(select_documents(docs), [])

    def test_keeps_only_first_document_per_group(self):
        from .risk_assessment import select_documents

        docs = [
            {"name": f"Проект контракта {i}.docx", "kind": "Проект контракта", "url": str(i)}
            for i in range(4)
        ]
        selected = select_documents(docs)
        self.assertEqual([d["name"] for d in selected], ["Проект контракта 0.docx"])

    def test_technical_spec_counts_as_description_group(self):
        from .risk_assessment import select_documents

        docs = [
            {"name": "Проект контракта.docx", "kind": "Проект контракта", "url": "a"},
            {"name": "Техническое задание.docx", "kind": "Техническое задание", "url": "b"},
        ]
        selected = select_documents(docs)
        self.assertEqual([d["name"] for d in selected], ["Проект контракта.docx", "Техническое задание.docx"])


class RiskAssessmentJsonParsingTests(TestCase):
    def test_strips_code_fence(self):
        from .risk_assessment import _json_from_model

        self.assertEqual(_json_from_model('```json\n{"a": 1}\n```'), {"a": 1})

    def test_recovers_from_extra_data_after_json(self):
        from .risk_assessment import _json_from_model

        self.assertEqual(_json_from_model('{"a": 1} что-то лишнее после'), {"a": 1})

    def test_raises_on_garbage(self):
        from .risk_assessment import RiskAssessmentError, _json_from_model

        with self.assertRaises(RiskAssessmentError):
            _json_from_model("совсем не похоже на json")


class RiskAssessmentRetryTests(TestCase):
    """assess() должен повторить запрос, если модель вернула не все обязательные ключи —
    живой тест 2026-09-16 показал, что так бывает даже без обрезания по max_tokens."""

    def test_complete_response_no_retry(self):
        from .risk_assessment import REQUIRED_KEYS, assess

        full = {key: "x" for key in REQUIRED_KEYS}
        with mock.patch("tender_selection.risk_assessment.call_gateway", return_value={"data": full, "usage": {}}) as call:
            result = assess("контекст")
        call.assert_called_once()
        self.assertEqual(result["data"], full)

    def test_incomplete_response_triggers_one_retry_and_keeps_better_result(self):
        from .risk_assessment import REQUIRED_KEYS, assess

        partial = {key: "x" for key in list(REQUIRED_KEYS)[:5]}
        full = {key: "x" for key in REQUIRED_KEYS}
        with mock.patch(
            "tender_selection.risk_assessment.call_gateway",
            side_effect=[{"data": partial, "usage": {}}, {"data": full, "usage": {}}],
        ) as call:
            result = assess("контекст")
        self.assertEqual(call.call_count, 2)
        self.assertEqual(result["data"], full)


class RiskAssessmentForTests(TestCase):
    """risk_assessment_for() — тот же паттерн кэширования/пометки попытки, что и
    у notification_for()/extras_for() (см. NotificationForTests, ExtrasForTests)."""

    def _tender(self, **kw):
        defaults = dict(purchase_number="1", object_info="x", title="T", law="fz44", last_pulled_at=timezone.now())
        defaults.update(kw)
        return FoundTender.objects.create(**defaults)

    def test_fz223_skipped(self):
        from .services import risk_assessment_for

        tender = self._tender(law="fz223")
        self.assertIsNone(risk_assessment_for(tender))

    def test_without_notification_returns_none(self):
        from .services import risk_assessment_for

        tender = self._tender()
        self.assertIsNone(risk_assessment_for(tender))

    def test_cache_is_used_without_recomputing(self):
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE, risk_assessment={"legal_risks": "уже оценено"})
        with mock.patch("tender_selection.risk_assessment.assess") as assess:
            result = risk_assessment_for(tender)
        assess.assert_not_called()
        self.assertEqual(result, {"legal_risks": "уже оценено"})

    def test_success_saves_result_and_used_documents(self):
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE)
        fake_result = {"data": {"legal_risks": "норм"}, "usage": {}}
        with mock.patch("tender_selection.services._fetch_doc_bytes", return_value=b"x"), \
             mock.patch("tender_selection.documents.extract_preview", return_value={"html": "<p>текст</p>"}), \
             mock.patch("tender_selection.risk_assessment.assess", return_value=fake_result) as assess:
            result = risk_assessment_for(tender)
        assess.assert_called_once()
        self.assertEqual(result, {"legal_risks": "норм"})
        tender.refresh_from_db()
        self.assertEqual(tender.risk_assessment, {"legal_risks": "норм"})
        self.assertEqual(tender.risk_assessment_docs, ["Описание объекта закупки.docx"])
        self.assertIsNotNone(tender.risk_checked_at)
        self.assertEqual(tender.risk_error, "")

    def test_force_recomputes_even_if_cached(self):
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE, risk_assessment={"legal_risks": "старое"})
        fake_result = {"data": {"legal_risks": "новое"}, "usage": {}}
        with mock.patch("tender_selection.services._fetch_doc_bytes", return_value=b"x"), \
             mock.patch("tender_selection.documents.extract_preview", return_value={"html": "<p>текст</p>"}), \
             mock.patch("tender_selection.risk_assessment.assess", return_value=fake_result):
            result = risk_assessment_for(tender, force=True)
        self.assertEqual(result, {"legal_risks": "новое"})

    def test_no_relevant_documents_records_error_without_calling_gateway(self):
        from .services import risk_assessment_for

        tender = self._tender(notification_raw={"doc_type": "epNotificationEF2020", "source": {}})
        with mock.patch("tender_selection.risk_assessment.assess") as assess:
            result = risk_assessment_for(tender)
        assess.assert_not_called()
        self.assertIsNone(result)
        tender.refresh_from_db()
        self.assertIn("документ", tender.risk_error)
        self.assertIsNotNone(tender.risk_checked_at)

    def test_all_documents_failing_falls_back_to_card_only_assessment(self):
        """Аварийный режим: если ни один документ не скачался (типично — локальная
        сеть не видит ЕИС), делаем оценку по одной сводке извещения из
        build_context(), явно помечая её degraded — чтобы не путать с полноценной
        оценкой по документам, а не молча пропускать запрос совсем."""
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE)
        fake_result = {"data": {"legal_risks": "по извещению"}, "usage": {}}
        with mock.patch("tender_selection.services._fetch_doc_bytes", side_effect=Exception("сеть недоступна")), \
             mock.patch("tender_selection.risk_assessment.assess", return_value=fake_result) as assess:
            result = risk_assessment_for(tender)
        assess.assert_called_once()
        self.assertEqual(result["legal_risks"], "по извещению")
        self.assertTrue(result["degraded"])
        tender.refresh_from_db()
        self.assertEqual(tender.risk_assessment_docs, [])
        self.assertEqual(tender.risk_error, "")

    def test_all_documents_failing_and_fallback_gateway_call_also_fails(self):
        from .risk_assessment import RiskAssessmentError
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE)
        with mock.patch("tender_selection.services._fetch_doc_bytes", side_effect=Exception("сеть недоступна")), \
             mock.patch("tender_selection.risk_assessment.assess", side_effect=RiskAssessmentError("нет ключа")):
            result = risk_assessment_for(tender)
        self.assertIsNone(result)
        tender.refresh_from_db()
        self.assertIn("нет ключа", tender.risk_error)

    def test_gateway_failure_recorded_not_raised(self):
        from .risk_assessment import RiskAssessmentError
        from .services import risk_assessment_for

        tender = self._tender(notification_raw=NOTIFICATION_FIXTURE)
        with mock.patch("tender_selection.services._fetch_doc_bytes", return_value=b"x"), \
             mock.patch("tender_selection.documents.extract_preview", return_value={"html": "<p>текст</p>"}), \
             mock.patch("tender_selection.risk_assessment.assess", side_effect=RiskAssessmentError("нет ключа")):
            result = risk_assessment_for(tender)
        self.assertIsNone(result)
        tender.refresh_from_db()
        self.assertEqual(tender.risk_error, "нет ключа")
        self.assertIsNotNone(tender.risk_checked_at)


class RiskStatusViewTests(TestCase):
    """Оценка рисков не должна блокировать открытие карточки (может идти десятки
    секунд — сеть до ЕИС) — первый расчёт уходит в JS-подгружаемый блок risk_status."""

    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("riskadmin", "risk@e.ru", "pw")
        self.client.force_login(self.admin)
        self.tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44",
            notification_raw=NOTIFICATION_FIXTURE, last_pulled_at=timezone.now(),
            review=FoundTender.INTERESTING,
        )

    def test_detail_page_shows_spinner_placeholder_without_calling_gateway(self):
        # "data-risk-loading" встречается и как JS-селектор в инлайн-скрипте, поэтому
        # SSR-состояние проверяем через context, а не грепом по HTML-тексту страницы.
        with mock.patch("tender_selection.risk_assessment.assess") as assess:
            resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        assess.assert_not_called()
        self.assertTrue(resp.context["risk_needs_fetch"])
        self.assertContains(resp, "data-risk-url")

    def test_cached_result_renders_inline_without_spinner(self):
        self.tender.risk_assessment = {"legal_risks": "уже оценено"}
        self.tender.risk_checked_at = timezone.now()
        self.tender.save(update_fields=["risk_assessment", "risk_checked_at"])
        resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertFalse(resp.context["risk_needs_fetch"])
        self.assertContains(resp, "уже оценено")

    def test_endpoint_computes_and_returns_html(self):
        fake_result = {"data": {"legal_risks": "норм"}, "usage": {}}
        with mock.patch("tender_selection.services._fetch_doc_bytes", return_value=b"x"), \
             mock.patch("tender_selection.documents.extract_preview", return_value={"html": "<p>текст</p>"}), \
             mock.patch("tender_selection.risk_assessment.assess", return_value=fake_result):
            resp = self.client.get(reverse("tender_selection:risk_status", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("норм", resp.json()["html"])
        self.tender.refresh_from_db()
        self.assertEqual(self.tender.risk_assessment, {"legal_risks": "норм"})

    def test_endpoint_respects_refresh_param(self):
        self.tender.risk_assessment = {"legal_risks": "старое"}
        self.tender.risk_checked_at = timezone.now()
        self.tender.save(update_fields=["risk_assessment", "risk_checked_at"])
        fake_result = {"data": {"legal_risks": "новое"}, "usage": {}}
        with mock.patch("tender_selection.services._fetch_doc_bytes", return_value=b"x"), \
             mock.patch("tender_selection.documents.extract_preview", return_value={"html": "<p>текст</p>"}), \
             mock.patch("tender_selection.risk_assessment.assess", return_value=fake_result) as assess:
            resp = self.client.get(reverse("tender_selection:risk_status", args=[self.tender.pk]) + "?refresh=1")
        assess.assert_called_once()
        self.assertIn("новое", resp.json()["html"])


class PriceStatsCollectorTests(TestCase):
    CONTRACT = {
        "purchase_number": "0111",
        "reg_num": "R1",
        "price": "168300.00",
        "okpd2": [],
        "ktru": ["17.23.13.196-00000001"],
        "region": 77,
        "subject": "Поставка канцелярских товаров",
        "suppliers": ["7712345678"],
        "published_at": "2026-09-01T10:00:00",
        "exe_start": "2026-09-05",
    }

    def test_category_from_codes(self):
        from tender_selection.stats import category_for_codes
        self.assertEqual(category_for_codes([], ["17.23.13.196-00000001"]), "17.23")
        self.assertEqual(category_for_codes(["32.99.53.190"], []), "32.99")
        self.assertEqual(category_for_codes(["99.00.00"], []), "")

    def test_collect_creates_row_then_fills_nmck_and_discount(self):
        from tender_selection.models import ContractStat
        from tender_selection.stats import collect_price_stats
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter([self.CONTRACT])), \
             mock.patch.object(gosplan, "fetch_purchase", return_value={"max_price": "210375.00"}), \
             mock.patch("tender_selection.stats.time.sleep"):
            run = collect_price_stats(since_days=30, catalog_requests=4, nmck_requests=6, categories=["17.23"])
        self.assertTrue(run.ok)
        self.assertEqual(run.created_count, 1)
        self.assertEqual(run.filled_count, 1)
        stat = ContractStat.objects.get(purchase_number="0111")
        self.assertEqual(stat.category, "17.23")
        self.assertEqual(str(stat.final_price), "168300.00")
        self.assertEqual(str(stat.nmck), "210375.00")
        self.assertEqual(str(stat.discount_pct), "20.0")
        self.assertEqual(stat.winner_inn, "7712345678")
        self.assertTrue(stat.nmck_checked)

    def test_second_run_does_not_refetch_nmck(self):
        from tender_selection.stats import collect_price_stats
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter([self.CONTRACT])), \
             mock.patch.object(gosplan, "fetch_purchase", return_value={"max_price": "210375.00"}), \
             mock.patch("tender_selection.stats.time.sleep"):
            collect_price_stats(catalog_requests=4, nmck_requests=6, categories=["17.23"])
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter([self.CONTRACT])), \
             mock.patch.object(gosplan, "fetch_purchase") as fp, \
             mock.patch("tender_selection.stats.time.sleep"):
            run2 = collect_price_stats(catalog_requests=4, nmck_requests=6, categories=["17.23"])
        fp.assert_not_called()
        self.assertEqual(run2.created_count, 0)

    def test_nmck_fill_respects_request_budget(self):
        from tender_selection.stats import collect_price_stats
        many = [dict(self.CONTRACT, purchase_number=f"n{i}", reg_num=f"R{i}") for i in range(8)]
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter(many)), \
             mock.patch.object(gosplan, "fetch_purchase", return_value={"max_price": "1000.00"}) as fp, \
             mock.patch("tender_selection.stats.time.sleep"):
            collect_price_stats(catalog_requests=4, nmck_requests=3, categories=["17.23"])
        self.assertEqual(fp.call_count, 3)

    def test_api_error_recorded_not_raised(self):
        from tender_selection.stats import collect_price_stats
        with mock.patch.object(gosplan, "iter_contracts", side_effect=gosplan.GosplanError("HTTP 429")), \
             mock.patch("tender_selection.stats.time.sleep"):
            run = collect_price_stats(catalog_requests=4, nmck_requests=6, categories=["17.23"])
        self.assertFalse(run.ok)
        self.assertIn("429", run.error)

    def test_multi_contract_purchase_marked_shared_and_skipped(self):
        from tender_selection.models import ContractStat
        from tender_selection.stats import collect_price_stats
        two = [
            dict(self.CONTRACT, purchase_number="P1", reg_num="RA", price="100000.00"),
            dict(self.CONTRACT, purchase_number="P1", reg_num="RB", price="300000.00"),
        ]
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter(two)), \
             mock.patch.object(gosplan, "fetch_purchase", return_value={"max_price": "1000000.00"}) as fp, \
             mock.patch("tender_selection.stats.time.sleep"):
            collect_price_stats(catalog_requests=4, nmck_requests=6, categories=["17.23"])
        rows = list(ContractStat.objects.filter(purchase_number="P1"))
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r.shared_purchase for r in rows))
        self.assertTrue(all(r.discount_pct is None for r in rows))
        fp.assert_not_called()

    def test_implausible_discount_dropped_but_nmck_kept(self):
        from tender_selection.models import ContractStat
        from tender_selection.stats import collect_price_stats
        with mock.patch.object(gosplan, "iter_contracts", return_value=iter([self.CONTRACT])), \
             mock.patch.object(gosplan, "fetch_purchase", return_value={"max_price": "5000000.00"}), \
             mock.patch("tender_selection.stats.time.sleep"):
            collect_price_stats(catalog_requests=4, nmck_requests=6, categories=["17.23"])
        stat = ContractStat.objects.get(purchase_number="0111")
        self.assertEqual(str(stat.nmck), "5000000.00")
        self.assertIsNone(stat.discount_pct)  # 96.6% — артефакт, отброшено
        self.assertTrue(stat.nmck_checked)


class PriceStatsCardTests(TestCase):
    def _seed_market(self, discounts, *, cat="32.99", nmck="500000", region=77,
                      subject="Поставка ежедневников"):
        import datetime
        from decimal import Decimal as D
        from tender_selection.models import ContractStat
        for i, d in enumerate(discounts):
            ContractStat.objects.create(
                law="fz44", purchase_number=f"p{i}", contract_reg_num=f"r{i}",
                category=cat, region=region, subject=f"{subject} {i}",
                nmck=D(nmck), final_price=D(nmck) * (100 - d) // 100,
                discount_pct=D(d), nmck_checked=True, contract_date=datetime.date(2026, 9, 1),
            )

    def _tender(self, **kw):
        from decimal import Decimal as D
        defaults = dict(purchase_number="X", law="fz44", object_info="x",
                        title="Поставка ежедневников с логотипом", okpd2=["32.99.11"],
                        max_price=D("500000"), region=77, last_pulled_at=timezone.now())
        defaults.update(kw)
        return FoundTender.objects.create(**defaults)

    def test_none_when_no_matching_words(self):
        from tender_selection.stats import price_stats_for
        self._seed_market((20, 30, 40), subject="Ремонт кровли гаража")
        self.assertIsNone(price_stats_for(self._tender()))

    def test_none_when_too_few_samples(self):
        from tender_selection.stats import price_stats_for
        self._seed_market((20, 30))
        self.assertIsNone(price_stats_for(self._tender()))

    def test_none_for_fz223(self):
        from tender_selection.stats import price_stats_for
        self._seed_market((10, 20, 30, 40, 50, 55))
        self.assertIsNone(price_stats_for(self._tender(law="fz223")))

    def test_aggregates_median_range_examples(self):
        from tender_selection.stats import price_stats_for
        self._seed_market((10, 20, 30, 40, 50, 55))
        s = price_stats_for(self._tender())
        self.assertEqual(s["count"], 6)
        self.assertEqual(s["own_count"], 0)
        self.assertEqual(s["market_count"], 6)
        self.assertEqual(s["median"], 35)               # median(10,20,30,40,50,55)
        self.assertEqual(s["suggested_reduction"], 35)
        self.assertEqual(s["same_region"], 6)
        self.assertEqual(len(s["examples"]), 6)

    def test_suggested_reduction_clamped(self):
        from tender_selection.stats import price_stats_for
        self._seed_market((70, 72, 75, 78, 80, 80))      # median 76.5 -> clamp to 60
        self.assertEqual(price_stats_for(self._tender())["suggested_reduction"], 60)

    def test_own_history_listed_before_market_and_preferred(self):
        """Своя история — сначала, рынок — только чтобы добрать до целевого числа."""
        from decimal import Decimal as D
        from django.contrib.auth import get_user_model
        from tender_selection.stats import price_stats_for
        from tenders.models import TenderEstimate, TenderLine

        self._seed_market((10, 20, 30, 40, 50))
        owner = get_user_model().objects.create_user("owner1", "o1@e.ru", "p")
        est = TenderEstimate.objects.create(
            owner=owner, tender_number="OWN1", name="Расчёт по ежедневникам",
            status=TenderEstimate.WON, actual_reduction_percent=D("22.00"),
            outcome_checked_at=timezone.now(),
        )
        TenderLine.objects.create(estimate=est, name="Ежедневники с тиснением", quantity=1, nmck_unit=1)

        s = price_stats_for(self._tender())
        self.assertEqual(s["own_count"], 1)
        self.assertEqual(s["market_count"], 5)
        self.assertEqual(s["examples"][0]["source"], "own")
        self.assertEqual(s["examples"][0]["discount_pct"], D("22.00"))

    def test_region_tiebreak_when_word_overlap_ties(self):
        from tender_selection.stats import price_stats_for, _TARGET_COUNT
        import datetime
        from decimal import Decimal as D
        from tender_selection.models import ContractStat
        for i in range(_TARGET_COUNT + 1):
            ContractStat.objects.create(
                law="fz44", purchase_number=f"tb{i}", contract_reg_num=f"tbr{i}", category="32.99",
                region=77 if i < 3 else 999, subject="Поставка ежедневников",
                nmck=D("500000"), final_price=D("400000"), discount_pct=D(10 + i),
                nmck_checked=True, contract_date=datetime.date(2026, 9, 1),
            )
        s = price_stats_for(self._tender())
        self.assertEqual(s["count"], _TARGET_COUNT)
        same_region = [row for row in s["examples"] if row["region"] == 77]
        self.assertEqual(len(same_region), 3)  # ни один "свой регион" не вытеснен при равном совпадении слов

    def test_detail_view_shows_section(self):
        from django.contrib.auth import get_user_model
        for name in ("fetch_clarifications", "fetch_complaints"):
            p = mock.patch.object(gosplan, name, return_value=[]); p.start(); self.addCleanup(p.stop)
        self._seed_market((15, 25, 35, 45, 50, 55))
        tender = self._tender(review=FoundTender.INTERESTING)
        self.client.force_login(get_user_model().objects.create_superuser("a", "a@e.ru", "p"))
        with mock.patch.object(gosplan, "fetch_notification", side_effect=gosplan.GosplanError("x")):
            resp = self.client.get(reverse("tender_selection:detail", args=[tender.pk]))
        self.assertContains(resp, "Снижение цены на похожих закупках")
        self.assertContains(resp, "подставим снижение")


class PriceStatsRelevanceTests(TestCase):
    def _row(self, i, discount, subject, **kw):
        import datetime
        from decimal import Decimal as D
        from tender_selection.models import ContractStat
        data = dict(
            law="fz44", purchase_number=f"p{i}", contract_reg_num=f"r{i}", category="32.99",
            region=1, subject=subject, nmck=D("500000"), final_price=D("400000"),
            discount_pct=D(discount), nmck_checked=True, contract_date=datetime.date(2026, 9, 1),
        )
        data.update(kw)
        return ContractStat.objects.create(**data)

    def _tender(self, **kw):
        from decimal import Decimal as D
        base = dict(
            purchase_number="T", law="fz44", object_info="x", title="Поставка ежедневников с логотипом",
            okpd2=["32.99.12.110"], max_price=D("500000"), region=77, customer_inn="7700000000",
            last_pulled_at=timezone.now(),
        )
        base.update(kw)
        return FoundTender.objects.create(**base)

    def test_ignores_same_category_different_product(self):
        from tender_selection.stats import price_stats_for
        for i in range(5):
            self._row(i, 30 + i, "Поставка ежедневников")
        for i in range(5, 15):  # та же категория ОКПД2, но другой товар — не должно попасть в выборку
            self._row(i, 4, "Поставка сувенирной продукции")
        s = price_stats_for(self._tender())
        self.assertEqual(s["count"], 5)
        self.assertEqual(s["median"], 32)

    def test_none_when_category_matches_but_product_does_not(self):
        from tender_selection.stats import price_stats_for
        for i in range(8):
            self._row(i, 20 + i, "Поставка сувенирной продукции")
        self.assertIsNone(price_stats_for(self._tender()))

    def test_own_archive_matches_by_line_item_not_generic_estimate_name(self):
        """Половина тендеров называется одинаково расплывчато ('поставка
        полиграфии'), а наполнение разное — сравнение должно смотреть на
        товарные позиции расчёта, а не только на его общее название."""
        from decimal import Decimal as D
        from django.contrib.auth import get_user_model
        from tender_selection.stats import price_stats_for
        from tenders.models import TenderEstimate, TenderLine

        owner = get_user_model().objects.create_user("owner2", "o2@e.ru", "p")
        matching = TenderEstimate.objects.create(
            owner=owner, tender_number="OWN2", name="Поставка полиграфии",
            status=TenderEstimate.WON, actual_reduction_percent=D("18.00"),
            outcome_checked_at=timezone.now(),
        )
        TenderLine.objects.create(estimate=matching, name="Ежедневники с тиснением", quantity=1, nmck_unit=1)
        unrelated = TenderEstimate.objects.create(
            owner=owner, tender_number="OWN3", name="Поставка полиграфии",
            status=TenderEstimate.LOST, actual_reduction_percent=D("5.00"),
            outcome_checked_at=timezone.now(),
        )
        TenderLine.objects.create(estimate=unrelated, name="Пакеты бумажные", quantity=1, nmck_unit=1)
        self._row(90, 25, "Поставка ежедневников")  # + пара рыночных, чтобы набрать минимум выборки
        self._row(91, 27, "Поставка ежедневников")

        s = price_stats_for(self._tender())
        self.assertEqual(s["own_count"], 1)
        self.assertEqual(s["examples"][0]["discount_pct"], D("18.00"))

    def test_keywords_ignore_boilerplate(self):
        from tender_selection.stats import _tender_keywords
        t = self._tender(title="Запрос котировок в электронной форме на поставку товаров для нужд учреждения")
        self.assertEqual(_tender_keywords(t), set())


class PushWithStatsTests(TestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        self.admin = get_user_model().objects.create_superuser("a", "a@e.ru", "p")
        for name in ("fetch_clarifications", "fetch_complaints"):
            p = mock.patch.object(gosplan, name, return_value=[]); p.start(); self.addCleanup(p.stop)

    def test_push_uses_suggested_reduction_and_stashes_summary(self):
        import datetime
        from decimal import Decimal as D
        from tender_selection.models import ContractStat
        from tender_selection.services import push_to_estimate
        for i, d in enumerate((20, 30, 30, 40, 45, 50)):  # median 35
            ContractStat.objects.create(
                law="fz44", purchase_number=f"c{i}", contract_reg_num=f"cr{i}", category="32.99",
                subject=f"Баннер {i}",  # тот же товар, что в позиции извещения (NOTIFICATION_FIXTURE)
                nmck=D("500000"), final_price=D("300000"), discount_pct=D(d),
                nmck_checked=True, contract_date=datetime.date(2026, 9, 1),
            )
        tender = FoundTender.objects.create(
            purchase_number="Z", law="fz44", object_info="x", title="Баннер", max_price=D("500000"),
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )
        est_id = push_to_estimate(tender, self.admin)
        from tenders.models import TenderEstimate
        est = TenderEstimate.objects.get(pk=est_id)
        self.assertEqual(str(est.reduction_percent), "35.00")
        self.assertEqual(est.summary_snapshot["price_stats"]["median"], 35)

    def test_push_without_stats_keeps_default_30(self):
        from decimal import Decimal as D
        from tender_selection.services import push_to_estimate
        tender = FoundTender.objects.create(
            purchase_number="Z", law="fz44", object_info="x", title="T", max_price=D("500000"),
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )
        est_id = push_to_estimate(tender, self.admin)
        from tenders.models import TenderEstimate
        est = TenderEstimate.objects.get(pk=est_id)
        self.assertEqual(str(est.reduction_percent), "30.00")
        self.assertNotIn("price_stats", est.summary_snapshot)


class PullCommandTests(TestCase):
    def test_non_loop_runs_one_pull(self):
        from django.core.management import call_command
        with mock.patch("tender_selection.management.commands.pull_tenders.run_pull") as rp:
            rp.return_value = mock.Mock(id=1, requests_made=1, records_received=0, created_count=0,
                                       updated_count=0, duration_seconds=1.0, ok=True, error="")
            call_command("pull_tenders", "--targeted")
        rp.assert_called_once()
        self.assertIsNone(rp.call_args.kwargs["classifiers"])


class IterPurchasesThrottleTests(TestCase):
    def test_stops_at_max_requests_and_sleeps_between(self):
        pages = [[{"purchase_number": str(i)} for i in range(100)] for _ in range(5)]
        calls = []

        def fake_fetch(params, law="fz44"):
            calls.append(params["skip"])
            return pages[len(calls) - 1]

        with mock.patch.object(gosplan, "fetch_page", side_effect=fake_fetch), \
             mock.patch.object(gosplan.time, "sleep") as sleep:
            got = list(gosplan.iter_purchases(params={}, max_requests=2))

        self.assertEqual(len(calls), 2)
        self.assertEqual(len(got), 200)
        self.assertEqual(sleep.call_count, 1)  # спит только между 1-м и 2-м запросом


class FilteringTests(TestCase):
    def test_plus_words_any_match_passes(self):
        inc = parse_terms("сувенир, поло, кружк")
        self.assertEqual(match_title("Поставка сувенирной продукции", inc, [])[0], True)
        self.assertEqual(match_title("Поставка бетона", inc, [])[0], False)

    def test_empty_plus_words_passes_everything(self):
        self.assertEqual(match_title("что угодно", [], [])[0], True)

    def test_minus_word_hides(self):
        inc = parse_terms("сувенир")
        exc = parse_terms("медал")  # подстрока, ловит «медали» и «медаль»
        passes, _ = match_title("Сувенирные медали", inc, exc)
        self.assertFalse(passes)

    def test_plus_inside_entry_needs_all_parts(self):
        inc = parse_terms("живых+цветов")
        self.assertFalse(match_title("Поставка живых кроликов", inc, [])[0])
        self.assertTrue(match_title("Букеты из живых цветов", inc, [])[0])

    def test_case_and_yo_insensitive(self):
        inc = parse_terms("ЁЛка")
        self.assertTrue(match_title("Поставка елочных игрушек", parse_terms("елочн"), [])[0])
        self.assertTrue(match_title("Новогодняя ЁЛКА", inc, [])[0])


class FormattingTests(TestCase):
    def test_rub_grouped_digits(self):
        from .templatetags.tender_selection_extras import rub
        nb = "\N{NO-BREAK SPACE}"
        self.assertEqual(rub(23966770), f"23{nb}966{nb}770{nb}₽")
        self.assertEqual(rub(1300000), f"1{nb}300{nb}000{nb}₽")
        self.assertEqual(rub(90000), f"90{nb}000{nb}₽")
        self.assertEqual(rub(None), "—")

    def test_region_name(self):
        from .regions import region_name
        self.assertEqual(region_name(78), "Санкт-Петербург")
        self.assertEqual(region_name(38), "Иркутская обл.")
        self.assertEqual(region_name(999), "регион 999")

    def test_parse_dt_treats_naive_as_utc(self):
        from .services import _parse_dt
        dt = _parse_dt("2026-09-15T06:00:00")
        self.assertEqual(dt.utcoffset().total_seconds(), 0)
        self.assertEqual(dt.astimezone(__import__("zoneinfo").ZoneInfo("Europe/Moscow")).hour, 9)


class SettingsViewTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.client.force_login(self.admin)

    def test_save_settings_persists_and_filters_list(self):
        for i, title in enumerate(["Поставка сувенирной продукции", "Поставка щебня"]):
            FoundTender.objects.create(
                purchase_number=str(i), object_info=title, title=title,
                max_price=500000, last_pulled_at=timezone.now(),
            )
        self.client.post(reverse("tender_selection:settings"), {
            "include_words": "сувенир", "exclude_words": "", "min_price": "300000", "window_days": "7",
        })
        s = FilterSettings.load()
        self.assertEqual(s.include_words, "сувенир")
        resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertContains(resp, "сувенирной")
        self.assertNotContains(resp, "щебня")
        self.assertContains(resp, "скрыто 1")

    def test_min_price_filters_list(self):
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Дорогая кружка", max_price=100000, last_pulled_at=timezone.now()
        )
        FilterSettings.objects.update_or_create(pk=1, defaults={"min_price": 300000})
        resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertNotContains(resp, "Дорогая кружка")
        self.assertContains(resp, "Ничего не подходит под фильтр")

    def test_inc_exc_query_override_without_touching_settings(self):
        FilterSettings.objects.update_or_create(pk=1, defaults={"include_words": "сувенир", "min_price": 0})
        for t in ("Поставка сувениров", "Поставка бланков строгой отчётности", "Поставка щебня"):
            FoundTender.objects.create(purchase_number=t[:20], object_info=t, title=t, last_pulled_at=timezone.now())
        # override: only "бланк" passes now
        resp = self.client.get(reverse("tender_selection:list") + "?view=list&inc=бланк")
        self.assertContains(resp, "бланков")
        self.assertNotContains(resp, "сувениров")
        # saved settings untouched
        self.assertEqual(FilterSettings.load().include_words, "сувенир")

    def test_save_words_updates_only_words(self):
        FilterSettings.objects.update_or_create(pk=1, defaults={
            "include_words": "старое", "exclude_words": "", "min_price": 300000,
        })
        self.client.post(reverse("tender_selection:save_words"), {"inc": "бланк, конверт", "exc": "б/у"})
        s = FilterSettings.load()
        self.assertEqual(s.include_words, "бланк, конверт")
        self.assertEqual(s.exclude_words, "б/у")
        self.assertEqual(s.min_price, 300000)  # прочие настройки не тронуты

    def test_sort_by_deadline(self):
        import datetime
        near = timezone.now() + datetime.timedelta(days=2)
        far = timezone.now() + datetime.timedelta(days=20)
        FoundTender.objects.create(purchase_number="near", object_info="A", title="A", collecting_finished_at=near, last_pulled_at=timezone.now())
        FoundTender.objects.create(purchase_number="far", object_info="B", title="B", collecting_finished_at=far, last_pulled_at=timezone.now())
        resp = self.client.get(reverse("tender_selection:list") + "?view=list&sort=deadline&inc=&exc=")
        body = resp.content.decode()
        self.assertLess(body.index("№ near"), body.index("№ far"))


NOTIFICATION_FIXTURE = {
    "doc_type": "epNotificationEF2020",
    "source": {
        "commonInfo": {
            "href": "https://zakupki.gov.ru/epz/order/notice/ea20/view/common-info.html?regNumber=1",
            "purchaseObjectInfo": "Поставка сувенирной продукции",
            "placingWay": {"name": "Электронный аукцион"},
            "ETP": {"name": "Сбербанк-АСТ", "url": "http://sberbank-ast.ru"},
        },
        "purchaseResponsibleInfo": {
            "responsibleOrgInfo": {
                "fullName": "ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ УЧРЕЖДЕНИЕ \"ПРИМЕР\"",
                "shortName": "ГБУ ПРИМЕР",
                "INN": "7700000000",
                "KPP": "770000000",
                "factAddress": "г Москва, ул Тестовая, 1",
            },
            "responsibleInfo": {
                "contactPersonInfo": {"lastName": "Иванов", "firstName": "Иван"},
                "contactEMail": "z@example.ru",
                "contactPhone": "7-495-0000000",
            },
        },
        "attachmentsInfo": {
            "attachmentInfo": [
                {
                    "fileName": "Описание объекта закупки.docx",
                    "fileSize": "8379349",
                    "docKindInfo": {"name": "Описание объекта закупки"},
                    "url": "https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=A",
                },
                {"fileName": "злая ссылка", "url": "https://evil.example/x"},
            ]
        },
        "notificationInfo": {
            "procedureInfo": {
                "collectingInfo": {"startDT": "2026-09-07T18:00:00+03:00", "endDT": "2026-09-15T09:00:00+03:00"},
                "biddingDate": "2026-09-16+03:00",
            },
            "contractConditionsInfo": {"maxPriceInfo": {"maxPrice": "1105700.00", "currency": {"name": "РОССИЙСКИЙ РУБЛЬ"}}},
            "customerRequirementsInfo": {
                "customerRequirementInfo": {
                    "applicationGuarantee": {"amount": "11057.00", "part": "1.0"},
                    "contractGuarantee": {"part": "5.0"},
                    "contractConditionsInfo": {
                        "deliveryPlacesInfo": {"byGARInfo": {"GARInfo": {"GARAddress": "г. Москва, ул. Пика, д. 4"}}},
                        "bankSupportContractRequiredInfo": {"treasurySupportContractInfo": {"treasurySupportContractRequired": "true"}},
                    },
                }
            },
            "purchaseObjectsInfo": {
                "notDrugPurchaseObjectsInfo": {
                    "purchaseObject": {
                        "name": "Баннер",
                        "OKPD2": {
                            "OKPDCode": "32.99.53.190",
                            "OKPDName": "Изделия демонстрационные прочие",
                            "characteristics": {
                                "characteristicsUsingTextForm": {
                                    "name": "Описание",
                                    "values": {"value": {"qualityDescription": "Размер 200х200, баннерная ткань"}},
                                }
                            },
                        },
                        "OKEI": {"nationalCode": "шт"},
                        "price": "3700.00",
                        "quantity": {"value": "1.00000000000"},
                        "sum": "3700.00",
                    }
                }
            },
        },
    },
}


class NotificationParseTests(TestCase):
    def setUp(self):
        from .notification import parse_notification
        self.card = parse_notification(NOTIFICATION_FIXTURE)

    def test_customer_and_contacts(self):
        self.assertIn("ПРИМЕР", self.card["customer"]["name"])
        self.assertEqual(self.card["customer"]["inn"], "7700000000")
        self.assertEqual(self.card["customer"]["contact_person"], "Иванов Иван")

    def test_dates_and_money(self):
        self.assertEqual(self.card["dates"]["collect_end"].hour, 9)
        self.assertEqual(self.card["money"]["max_price"], "1105700.00")
        self.assertEqual(self.card["money"]["app_guarantee_part"], "1.0")
        self.assertTrue(self.card["money"]["treasury_support"])

    def test_single_item_becomes_list_with_characteristics(self):
        self.assertEqual(len(self.card["items"]), 1)
        item = self.card["items"][0]
        self.assertEqual(item["name"], "Баннер")
        self.assertEqual(item["code"], "32.99.53.190")
        self.assertEqual(item["characteristics"][0]["value"], "Размер 200х200, баннерная ткань")

    def test_only_zakupki_documents_kept(self):
        self.assertEqual(len(self.card["documents"]), 1)
        self.assertEqual(self.card["documents"][0]["size_kb"], 8183)
        self.assertTrue(self.card["documents"][0]["url"].startswith("https://zakupki.gov.ru/"))

    def test_delivery_address(self):
        self.assertIn("Пика", self.card["delivery_address"])


CLARIFICATION_FIXTURE = [
    {
        "doc_type": "epClarificationOfDocumentationEZ44",
        "published_at": "2026-09-05T10:00:00",
        "source": {
            "commonInfo": {
                "docNumber": "РД1",
                "href": "https://zakupki.gov.ru/epz/order/notice/ea44/view/clarifications.html?regNumber=1",
            },
            "clarificationInfo": {
                "questionText": "Просим уточнить требуемый размер флага и материал полотна для позиции 1.",
                "answerText": "Размер флага 90x135 см, материал — флажная сетка, плотность 115 г/кв.м.",
            },
            "extPrintFormInfo": {
                "url": "https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=CLR",
            },
        },
    },
    {"doc_type": "epExplanation", "published_at": "2026-09-04T09:00:00", "source": {}},
]

COMPLAINT_FIXTURE = [
    {
        "reg_number": "202600126646001665",
        "published_at": "2026-09-07T19:34:32",
        "updated_at": "2026-09-07T19:34:32",
        "purchase_number": "1",
        "object": "purchase",
        "region": 15,
        "docs": [{"doc_type": "complaint", "published_at": "2026-09-07T19:34:32"}],
    },
]


class ClarificationComplaintParseTests(TestCase):
    def test_clarification_extracts_question_and_answer(self):
        from .notification import parse_clarifications
        rows = parse_clarifications(CLARIFICATION_FIXTURE)
        self.assertEqual(len(rows), 2)
        first = rows[0]  # новее — от 05.09
        self.assertEqual(first["label"], "Разъяснение документации")
        self.assertIn("размер флага", first["question"])
        self.assertIn("90x135", first["answer"])
        self.assertTrue(first["print_url"].startswith("https://zakupki.gov.ru/"))
        self.assertTrue(first["href"].startswith("https://zakupki.gov.ru/"))

    def test_clarification_minimal_source_still_a_row(self):
        from .notification import parse_clarifications
        bare = parse_clarifications(CLARIFICATION_FIXTURE)[1]
        self.assertEqual(bare["label"], "Разъяснение")
        self.assertEqual(bare["question"], "")
        self.assertEqual(bare["answer"], "")
        self.assertIsNotNone(bare["published_at"])

    def test_complaints_metadata(self):
        from .notification import parse_complaints
        rows = parse_complaints(COMPLAINT_FIXTURE)
        self.assertEqual(rows[0]["kinds"], ["Жалоба"])
        self.assertEqual(rows[0]["object"], "закупка")
        self.assertEqual(rows[0]["reg_number"], "202600126646001665")


class ExtrasForTests(TestCase):
    def setUp(self):
        self.tender = FoundTender.objects.create(
            purchase_number="0345", law="fz44", object_info="x", title="T", last_pulled_at=timezone.now(),
        )

    def test_first_call_fetches_then_cache_serves_repeat(self):
        from .services import extras_for
        with mock.patch.object(gosplan, "fetch_clarifications", return_value=CLARIFICATION_FIXTURE) as fc, \
             mock.patch.object(gosplan, "fetch_complaints", return_value=COMPLAINT_FIXTURE):
            clar, comp = extras_for(self.tender)
        self.assertEqual(len(clar), 2)
        self.assertEqual(len(comp), 1)
        fc.assert_called_once()
        self.tender.refresh_from_db()
        self.assertIsNotNone(self.tender.extras_checked_at)
        with mock.patch.object(gosplan, "fetch_clarifications") as fc2, \
             mock.patch.object(gosplan, "fetch_complaints") as fk2:
            extras_for(self.tender)
        fc2.assert_not_called()
        fk2.assert_not_called()

    def test_force_refetches(self):
        from .services import extras_for
        with mock.patch.object(gosplan, "fetch_clarifications", return_value=[]), \
             mock.patch.object(gosplan, "fetch_complaints", return_value=[]):
            extras_for(self.tender)
        with mock.patch.object(gosplan, "fetch_clarifications", return_value=CLARIFICATION_FIXTURE) as fc, \
             mock.patch.object(gosplan, "fetch_complaints", return_value=[]):
            clar, _ = extras_for(self.tender, force=True)
        fc.assert_called_once()
        self.assertEqual(len(clar), 2)

    def test_fz223_skips_api(self):
        from .services import extras_for
        t = FoundTender.objects.create(purchase_number="9", law="fz223", object_info="x", last_pulled_at=timezone.now())
        with mock.patch.object(gosplan, "fetch_clarifications") as fc:
            self.assertEqual(extras_for(t), ([], []))
        fc.assert_not_called()

    def test_api_error_neither_raises_nor_caches(self):
        from .services import extras_for
        with mock.patch.object(gosplan, "fetch_clarifications", side_effect=gosplan.GosplanError("429")), \
             mock.patch.object(gosplan, "fetch_complaints", side_effect=gosplan.GosplanError("429")):
            self.assertEqual(extras_for(self.tender), ([], []))
        self.tender.refresh_from_db()
        self.assertIsNone(self.tender.extras_checked_at)


class ClarificationComplaintViewTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.client.force_login(self.admin)
        self.tender = FoundTender.objects.create(
            purchase_number="1", law="fz44", object_info="x", title="Флаги с логотипом",
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )

    def test_detail_renders_both_sections(self):
        with mock.patch.object(gosplan, "fetch_clarifications", return_value=CLARIFICATION_FIXTURE), \
             mock.patch.object(gosplan, "fetch_complaints", return_value=COMPLAINT_FIXTURE):
            resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertContains(resp, "Разъяснения")
        self.assertContains(resp, "90x135")
        self.assertContains(resp, "Жалобы в ФАС")
        self.assertContains(resp, "202600126646001665")

    def test_complaint_flag_shows_in_list_after_check(self):
        with mock.patch.object(gosplan, "fetch_clarifications", return_value=[]), \
             mock.patch.object(gosplan, "fetch_complaints", return_value=COMPLAINT_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertContains(resp, "ts-flag")


class DetailViewTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.client.force_login(self.admin)
        for name in ("fetch_clarifications", "fetch_complaints"):
            patcher = mock.patch.object(gosplan, name, return_value=[])
            patcher.start()
            self.addCleanup(patcher.stop)
        self.tender = FoundTender.objects.create(
            purchase_number="0345500000626000081", object_info="x", title="Сувенирка",
            customer_inn="7700000000", last_pulled_at=timezone.now(),
        )

    def test_fetches_and_renders_card(self):
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE) as fetch:
            resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Баннер")
        self.assertContains(resp, "Описание объекта закупки.docx")
        self.assertContains(resp, "ГБУ ПРИМЕР")
        fetch.assert_called_once()
        self.tender.refresh_from_db()
        self.assertTrue(self.tender.notification_raw)

    def test_second_open_uses_cache_no_request(self):
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        with mock.patch.object(gosplan, "fetch_notification") as fetch:
            resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        fetch.assert_not_called()
        self.assertEqual(resp.status_code, 200)

    def test_customer_name_backfilled_to_list(self):
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertTrue(Organization.objects.filter(inn="7700000000", name__icontains="ПРИМЕР").exists())

    def test_fetch_failure_renders_fallback(self):
        with mock.patch.object(gosplan, "fetch_notification", side_effect=gosplan.RateLimitError("429")):
            resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Не удалось загрузить извещение")


class OpenedAtTests(TestCase):
    """«Непрочитанные» тендеры выделяются в списке жирным (is-unread) — открытие
    карточки должно проставлять opened_at один раз и не трогать его повторно."""

    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("openadmin", "open@e.ru", "pw")
        self.client.force_login(self.admin)
        for name in ("fetch_clarifications", "fetch_complaints"):
            patcher = mock.patch.object(gosplan, name, return_value=[])
            patcher.start()
            self.addCleanup(patcher.stop)

    def test_first_open_sets_opened_at(self):
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44", last_pulled_at=timezone.now(),
        )
        self.assertIsNone(tender.opened_at)
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[tender.pk]))
        tender.refresh_from_db()
        self.assertIsNotNone(tender.opened_at)

    def test_second_open_does_not_change_opened_at(self):
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", title="T", law="fz44", last_pulled_at=timezone.now(),
        )
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[tender.pk]))
        tender.refresh_from_db()
        first = tender.opened_at
        with mock.patch.object(gosplan, "fetch_notification", return_value=NOTIFICATION_FIXTURE):
            self.client.get(reverse("tender_selection:detail", args=[tender.pk]))
        tender.refresh_from_db()
        self.assertEqual(tender.opened_at, first)

    def test_list_marks_unopened_tender_unread(self):
        # плюс/минус-слова (см. миграции 0013/0014) по умолчанию скрывают тендеры
        # с не подходящим по смыслу названием — ?all=1 показывает все, как и сама
        # ссылка «Показать все» в шаблоне.
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="ТендерОдинЕщёНеОткрыт", last_pulled_at=timezone.now(),
        )
        FoundTender.objects.create(
            purchase_number="2", object_info="x", title="ТендерДваУжеОткрыт", opened_at=timezone.now(),
            last_pulled_at=timezone.now(),
        )
        resp = self.client.get(reverse("tender_selection:list") + "?view=list&all=1")
        content = resp.content.decode()
        unread_pos = content.index("ТендерОдинЕщёНеОткрыт")
        read_pos = content.index("ТендерДваУжеОткрыт")
        self.assertIn("is-unread", content[max(0, unread_pos - 400):unread_pos])
        self.assertNotIn("is-unread", content[max(0, read_pos - 400):read_pos])


class AccessControlTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.user = User.objects.create_user("bob", "b@e.ru", "pw")

    def test_anonymous_redirected_to_login(self):
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/login", resp["Location"])

    def test_regular_user_forbidden(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertEqual(resp.status_code, 403)

    def test_superuser_sees_the_list(self):
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружки", last_pulled_at=timezone.now()
        )
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Кружки")

    def test_set_review_ajax(self):
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", last_pulled_at=timezone.now()
        )
        self.client.force_login(self.admin)
        resp = self.client.post(
            reverse("tender_selection:review", args=[tender.pk]),
            {"review": "interesting"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(resp.json()["review"], "interesting")
        tender.refresh_from_db()
        self.assertEqual(tender.review, FoundTender.INTERESTING)

    def test_set_review_rejects_unknown_value(self):
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", last_pulled_at=timezone.now()
        )
        self.client.force_login(self.admin)
        self.client.post(reverse("tender_selection:review", args=[tender.pk]), {"review": "bogus"})
        tender.refresh_from_db()
        self.assertEqual(tender.review, FoundTender.UNREVIEWED)

    def test_list_row_has_no_forward_button_only_hide(self):
        """Решение «вперёд» (На оценку рисков / В расчёт) принимается только на
        странице самого тендера — рано жать кнопку из списка, не открыв, что
        внутри. В списке остаётся только «×» скрыть, независимо от review."""
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружка", review=FoundTender.NOT_INTERESTING,
            last_pulled_at=timezone.now(),
        )
        FoundTender.objects.create(
            purchase_number="2", object_info="x", title="Блокнот", last_pulled_at=timezone.now(),
        )
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        self.assertNotContains(resp, "На оценку рисков")
        self.assertNotContains(resp, "В расчёт")
        self.assertContains(resp, 'title="Скрыть тендер"')

    def test_missing_notification_shows_warning_badge(self):
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружка", law="fz44",
            notification_checked_at=timezone.now(), last_pulled_at=timezone.now(),
        )  # notification_raw пуст, но попытка БЫЛА (checked_at стоит) — реальный сбой
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list") + "?view=list")
        # class="ts-flag--data" встречается только у самого <span> — не путать с
        # правилом .ts-flag--data в <style> того же шаблона.
        self.assertContains(resp, "ts-flag ts-flag--data")

    def test_never_checked_notification_has_no_warning_badge(self):
        """Свежевыгруженный тендер, извещение для которого ещё никогда не запрашивали
        (notification_checked_at пуст) — это не сбой API, а «ещё не проверяли», значок
        не должен гореть просто потому, что карточку никто не открывал (регрессия:
        раньше бейдж стоял абсолютно на всех новых тендерах)."""
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружка", law="fz44",
            last_pulled_at=timezone.now(),
        )  # notification_raw и notification_checked_at пусты по умолчанию
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertNotContains(resp, "ts-flag ts-flag--data")

    def test_loaded_notification_hides_warning_badge(self):
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружка", law="fz44",
            notification_raw={"source": {}}, last_pulled_at=timezone.now(),
        )
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertNotContains(resp, "ts-flag ts-flag--data")

    def test_fz223_without_notification_has_no_warning_badge(self):
        FoundTender.objects.create(
            purchase_number="1", object_info="x", title="Кружка", law="fz223",
            last_pulled_at=timezone.now(),
        )  # у 223-ФЗ нет разобранного извещения по конструкции — это не ошибка
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertNotContains(resp, "ts-flag ts-flag--data")

    def test_dismiss_hides_row(self):
        tender = FoundTender.objects.create(
            purchase_number="1", object_info="x", last_pulled_at=timezone.now()
        )
        self.client.force_login(self.admin)
        self.client.post(reverse("tender_selection:dismiss", args=[tender.pk]))
        tender.refresh_from_db()
        self.assertEqual(tender.status, FoundTender.DISMISSED)


class TenderViewerAccessTests(TestCase):
    """Менеджер (не суперюзер) может открыть карточку СВОЕГО тендера — того, что
    стало его просчётом — по ссылке «Открыть карточку →» со страницы просчёта, но
    не листать каталог подбора (tender_list остаётся суперюзер-only, см.
    AccessControlTests.test_regular_user_forbidden)."""

    def setUp(self):
        from tenders.models import TenderEstimate

        User = get_user_model()
        self.admin = User.objects.create_superuser("admin", "a@e.ru", "pw")
        self.manager = User.objects.create_user("mgr", "m@e.ru", "pw")
        self.other_manager = User.objects.create_user("other", "o@e.ru", "pw")
        self.estimate = TenderEstimate.objects.create(
            owner=self.manager, tender_number="1", name="Просчёт менеджера",
        )
        self.tender = FoundTender.objects.create(
            purchase_number="1", law="fz44", object_info="x", title="Сувенирка",
            status=FoundTender.PUSHED, pushed_estimate_id=self.estimate.pk,
            last_pulled_at=timezone.now(), notification_raw=NOTIFICATION_FIXTURE,
        )

    def test_owner_can_open_own_tender_detail(self):
        self.client.force_login(self.manager)
        resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 200)

    def test_owner_sees_link_back_to_estimate_not_catalog(self):
        # Не просто "содержит /tender-selection/" — этот префикс встречается и в
        # других ссылках на той же странице (просмотр/скачивание документов).
        # Проверяем сам текст ссылки-навигации наверху карточки.
        self.client.force_login(self.manager)
        resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertContains(resp, f'href="{reverse("tender_estimate", args=[self.estimate.pk])}" class="back-link">← назад к просчёту')
        self.assertNotContains(resp, "к списку")

    def test_superuser_still_sees_link_to_catalog(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertContains(resp, f'href="{reverse("tender_selection:list")}" class="back-link">← к списку')

    def test_other_manager_cannot_open_someone_elses_tender(self):
        self.client.force_login(self.other_manager)
        resp = self.client.get(reverse("tender_selection:detail", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 404)

    def test_manager_cannot_open_not_yet_pushed_tender(self):
        fresh = FoundTender.objects.create(
            purchase_number="2", law="fz44", object_info="x", title="Ещё не в расчёте",
            last_pulled_at=timezone.now(),
        )
        self.client.force_login(self.manager)
        resp = self.client.get(reverse("tender_selection:detail", args=[fresh.pk]))
        self.assertEqual(resp.status_code, 404)

    def test_manager_still_forbidden_from_catalog_list(self):
        self.client.force_login(self.manager)
        resp = self.client.get(reverse("tender_selection:list"))
        self.assertEqual(resp.status_code, 403)

    def test_manager_still_forbidden_from_admin_actions(self):
        self.client.force_login(self.manager)
        resp = self.client.post(reverse("tender_selection:dismiss", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 403)
        resp = self.client.post(reverse("tender_selection:review", args=[self.tender.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_owner_can_preview_own_tender_documents(self):
        from .documents import DocumentError
        from .eis_docs import EisDocsError

        self.client.force_login(self.manager)
        with mock.patch("tender_selection.services.fetch_document_via_eis", side_effect=EisDocsError("нет сети")), \
             mock.patch("tender_selection.services.fetch_document", side_effect=DocumentError("нет сети")):
            resp = self.client.get(reverse("tender_selection:doc_preview", args=[self.tender.pk, 0]))
        self.assertEqual(resp.status_code, 200)  # прошёл контроль доступа, дошёл до бизнес-логики

    def test_other_manager_cannot_preview_documents(self):
        self.client.force_login(self.other_manager)
        resp = self.client.get(reverse("tender_selection:doc_preview", args=[self.tender.pk, 0]))
        self.assertEqual(resp.status_code, 404)
