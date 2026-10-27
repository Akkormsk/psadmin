import io
from pathlib import Path
import subprocess
from unittest.mock import patch

from django.test import SimpleTestCase
from docx import Document
from openpyxl import Workbook

from .documents import DocumentError, extract_preview
from .office import convert_office


class OfficeConversionTests(SimpleTestCase):
    def test_conversion_uses_isolated_profile_and_removes_temporary_files(self):
        paths = []

        def run(command, **kwargs):
            source = Path(command[-1])
            paths.append(source.parent)
            self.assertEqual(source.name, "document.docx")
            self.assertIn("--headless", command)
            self.assertTrue(any(arg.startswith("-env:UserInstallation=file:") for arg in command))
            self.assertEqual(kwargs["timeout"], 60)
            (source.parent / "output" / "document.pdf").write_bytes(b"%PDF-test")

        with patch("tender_selection.office.shutil.which", return_value="/usr/bin/libreoffice"), \
                patch("tender_selection.office.subprocess.run", side_effect=run):
            self.assertEqual(convert_office(b"office", "../../name.docx", "pdf"), b"%PDF-test")
        self.assertFalse(paths[0].exists())

    def test_timeout_is_explicit_failure(self):
        with patch("tender_selection.office.shutil.which", return_value="libreoffice"), \
                patch("tender_selection.office.subprocess.run", side_effect=subprocess.TimeoutExpired("office", 60)):
            with self.assertRaises(DocumentError):
                convert_office(b"office", "name.xlsx", "pdf")

    def test_missing_converter_is_explicit_failure(self):
        with patch("tender_selection.office.shutil.which", return_value=None):
            with self.assertRaises(DocumentError):
                convert_office(b"office", "name.doc", "docx")

    def test_legacy_word_uses_structured_reader_after_conversion(self):
        document = Document()
        document.add_paragraph("Русский текст без искажений")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Состав комплекта"
        output = io.BytesIO()
        document.save(output)
        with patch("tender_selection.office.convert_office", return_value=output.getvalue()) as converter:
            result = extract_preview(b"legacy", "source.doc")
        converter.assert_called_once_with(b"legacy", "source.doc", "docx")
        self.assertIn("Русский текст без искажений", result["html"])
        self.assertIn("Состав комплекта", result["html"])

    def test_legacy_excel_uses_structured_reader_after_conversion(self):
        workbook = Workbook()
        workbook.active.append(["Товар", 123])
        output = io.BytesIO()
        workbook.save(output)
        with patch("tender_selection.office.convert_office", return_value=output.getvalue()):
            result = extract_preview(b"legacy", "source.xls")
        self.assertIn("Товар", result["html"])
        self.assertIn("123", result["html"])

    def test_modern_word_parse_failure_uses_converter_fallback(self):
        with patch("tender_selection.documents._docx_html", side_effect=[ValueError("parse"), "<p>Текст</p>"]), \
                patch("tender_selection.office.convert_office", return_value=b"normalized") as convert:
            result = extract_preview(b"input", "source.docx")
        convert.assert_called_once_with(b"input", "source.docx", "docx")
        self.assertEqual(result["html"], "<p>Текст</p>")

    def test_excel_pdf_uses_whole_sheet_export(self):
        def run(command, **kwargs):
            output_format = command[command.index("--convert-to") + 1]
            self.assertIn('"SinglePageSheets"', output_format)
            (Path(command[-1]).parent / "output" / "document.pdf").write_bytes(b"%PDF-test")
        with patch("tender_selection.office.shutil.which", return_value="libreoffice"), \
                patch("tender_selection.office.subprocess.run", side_effect=run):
            convert_office(b"excel", "table.xlsx", "pdf")
